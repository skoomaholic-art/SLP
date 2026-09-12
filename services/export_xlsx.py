from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from services.live_evidence import event_is_live_broadcast
from services.time_logic import KZ_TIMEZONE, get_event_status, get_scheduled_datetimes


EXPORT_MODE_ALL = "all"
EXPORT_MODE_LIVE = "live"
EXPORT_MODE_FINISH_UPCOMING = "finish-upcoming"
EXPORT_MODES = {
    EXPORT_MODE_ALL,
    EXPORT_MODE_LIVE,
    EXPORT_MODE_FINISH_UPCOMING,
}

HEADERS = (
    "Дата",
    "Начало",
    "Окончание",
    "Статус эфира",
    "Вид спорта",
    "Турнир",
    "Событие",
    "Канал",
    "Источник",
    "Исходное название",
    "URL источника",
    "Прямой эфир",
    "Часовой пояс",
)


def _title(event: dict) -> str:
    return str(event.get("title") or event.get("raw_title") or "")


def _normalize_now(now: datetime | None) -> datetime:
    if now is None:
        return datetime.now(KZ_TIMEZONE)
    if now.tzinfo is None:
        return now.replace(tzinfo=KZ_TIMEZONE)
    return now.astimezone(KZ_TIMEZONE)


def filter_export_events(
    events: list[dict],
    *,
    mode: str = EXPORT_MODE_ALL,
    now: datetime | None = None,
) -> list[dict]:
    if mode not in EXPORT_MODES:
        raise ValueError(f"Неизвестный режим XLSX: {mode}")

    reference = _normalize_now(now)
    if mode == EXPORT_MODE_ALL:
        return list(events)
    if mode == EXPORT_MODE_LIVE:
        return [
            event for event in events
            if get_event_status(event, now=reference) == "live"
        ]
    return [
        event for event in events
        if get_event_status(event, now=reference) in {"finished", "upcoming"}
    ]


def build_schedule_xlsx(
    events: list[dict],
    *,
    mode: str = EXPORT_MODE_ALL,
    now: datetime | None = None,
) -> bytes:
    reference = _normalize_now(now)
    selected = filter_export_events(events, mode=mode, now=reference)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Расписание"
    sheet.append(HEADERS)

    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for event in selected:
        start, end = get_scheduled_datetimes(event)
        sheet.append(
            (
                start.strftime("%d.%m.%Y"),
                start.strftime("%H:%M"),
                end.strftime("%d.%m.%Y %H:%M"),
                get_event_status(event, now=reference).upper(),
                str(event.get("sport") or ""),
                str(event.get("tournament") or ""),
                _title(event),
                str(event.get("channel") or ""),
                str(event.get("source") or ""),
                str(event.get("raw_title") or ""),
                str(event.get("source_url") or ""),
                "Да" if event_is_live_broadcast(event) else "Нет",
                str(event.get("timezone") or "Asia/Almaty"),
            )
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    widths = (12, 10, 20, 16, 18, 30, 44, 24, 16, 48, 48, 14, 18)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    summary = workbook.create_sheet("Сводка")
    summary.append(("Показатель", "Значение"))
    for cell in summary[1]:
        cell.font = Font(bold=True)

    status_counts = {"live": 0, "upcoming": 0, "finished": 0}
    channels: dict[str, int] = {}
    sources: dict[str, int] = {}
    direct_count = 0

    for event in selected:
        status = get_event_status(event, now=reference)
        status_counts[status] = status_counts.get(status, 0) + 1
        if event_is_live_broadcast(event):
            direct_count += 1
        channel = str(event.get("channel") or "Не указан")
        source = str(event.get("source") or "Не указан")
        channels[channel] = channels.get(channel, 0) + 1
        sources[source] = sources.get(source, 0) + 1

    summary.append(("Режим выгрузки", mode))
    summary.append(("Точка расчёта", reference.isoformat(timespec="minutes")))
    summary.append(("Часовой пояс", "Asia/Almaty (UTC+5)"))
    summary.append(("Всего событий", len(selected)))
    summary.append(("LIVE (сейчас в эфире)", status_counts.get("live", 0)))
    summary.append(("UPCOMING", status_counts.get("upcoming", 0)))
    summary.append(("FINISHED", status_counts.get("finished", 0)))
    summary.append(("Прямой эфир = Да", direct_count))
    summary.append(())
    summary.append(("Канал", "Событий"))
    for channel, count in sorted(channels.items()):
        summary.append((channel, count))
    summary.append(())
    summary.append(("Источник", "Событий"))
    for source, count in sorted(sources.items()):
        summary.append((source, count))

    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 28

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
