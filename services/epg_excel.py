"""Import genuine Setanta/QSport .xlsx EPGs without treating ordinary reruns as LIVE.

The importer is independent of Gmail transport: downloaded attachments and
an authorized manual upload use exactly the same verified parser.
No raw workbook bytes or email credentials are stored in git or SQLite.
"""
from __future__ import annotations

import hashlib
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from io import BytesIO
from pathlib import PurePath
from zipfile import ZipFile, BadZipFile
import re
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from services.live_evidence import event_is_editorial_or_replay
from storage.database import SLPDatabase

KZ = ZoneInfo("Asia/Almaty")
MAX_WORKBOOK_BYTES = 6 * 1024 * 1024
MAX_SHEET_ROWS = 5000
MONTHS = {
    "января": 1, "февраля": 2, "марта": 3, "апреля": 4,
    "мая": 5, "июня": 6, "июля": 7, "августа": 8,
    "сентября": 9, "октября": 10, "ноября": 11, "декабря": 12,
}
MONTH_PATTERN = "|".join(MONTHS)
DAY_RE = re.compile(r"\s*(\d{1,2})\s+(" + MONTH_PATTERN + r")(?:\s+(\d{4}))?\s*", re.I)
LIVE_RE = re.compile(r"^\s*LIVE\s*[\.:,\-]\s*", re.I)
SPORTS = {
    "футбол": "Футбол", "футзал": "Футзал", "баскетбол": "Баскетбол",
    "волейбол": "Волейбол", "пляжный волейбол": "Пляжный волейбол",
    "теннис": "Теннис", "хоккей": "Хоккей", "мма": "ММА",
    "mma": "ММА", "бокс": "Бокс", "дзюдо": "Дзюдо",
    "мотоспорт": "Мотоспорт", "формула 1": "Формула-1",
    "формула-1": "Формула-1", "снукер": "Снукер",
    "фигурное катание": "Фигурное катание",
    "киберспорт": "Киберспорт", "гандбол": "Гандбол",
    "регби": "Регби", "велоспорт": "Велоспорт",
    "легкая атлетика": "Лёгкая атлетика",
    "лёгкая атлетика": "Лёгкая атлетика",
    "женіл атлетика": "Лёгкая атлетика",
    "жеңіл атлетика": "Лёгкая атлетика",
    "гимнастика": "Гимнастика",
    "таэквондо": "Таэквондо", "карате": "Карате",
    "борьба": "Борьба", "күрес": "Борьба",
    "бильярд": "Бильярд", "гольф": "Гольф",
    "лыжные гонки": "Лыжный спорт", "биатлон": "Биатлон",
}
CHANNELS = {
    "setanta1": "SETANTA SPORTS 1",
    "setanta2": "SETANTA SPORTS 2",
    "setantakz": "SETANTA SPORTS KZ",
    "qleague": "Q LEAGUE",
    "qarena": "Q ARENA",
    "qfootball": "Q FOOTBALL",
}
SOURCE_KEY = {channel: "email_epg_" + key for key, channel in CHANNELS.items()}


class InvalidEPG(ValueError):
    pass


@dataclass(frozen=True)
class ParsedEPG:
    channel: str
    filename: str
    content_hash: str
    scope_dates: tuple[str, ...]
    all_programmes: int
    events: tuple[dict[str, Any], ...]


def detect_channel(filename: str) -> str:
    name = re.sub(r"\s+", " ", filename.casefold())
    if "setanta" in name:
        if re.search(r"\bsports\s*1\b", name):
            return CHANNELS["setanta1"]
        if re.search(r"\bsports\s*2\b", name):
            return CHANNELS["setanta2"]
        if re.search(r"\b(qazaqstan|kazakhstan\s*kz|setanta\s*kz)\b", name):
            return CHANNELS["setantakz"]
    if re.search(r"\bq\s*sport\b|\bqsport\b", name):
        for marker, key in (
            ("league", "qleague"), ("arena", "qarena"),
            ("football", "qfootball"),
        ):
            if re.search(r"\b" + marker + r"\b", name):
                return CHANNELS[key]
    raise InvalidEPG("Не удалось подтвердить телеканал по названию файла")


def _year(filename: str, today: date) -> int:
    explicit = re.search(r"\b(20\d{2})\b", filename)
    if explicit:
        return int(explicit.group(1))
    # Example: 29.09.26 - 05.10.26 in a weekly supplier filename.
    short = re.search(r"\b\d{1,2}[.]\d{1,2}[.](\d{2})\b", filename)
    if short:
        return 2000 + int(short.group(1))
    return today.year


def _day_header(value: Any, year: int) -> date | None:
    if not isinstance(value, str):
        return None
    match = DAY_RE.fullmatch(value.strip())
    if not match:
        return None
    day, month, explicit_year = match.groups()
    try:
        return date(int(explicit_year) if explicit_year else year,
                    MONTHS[month.casefold()], int(day))
    except ValueError:
        return None


def _clock_minutes(value: Any) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, datetime):
        return value.hour * 60 + value.minute
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    if isinstance(value, timedelta):
        return round(value.total_seconds() / 60)
    if isinstance(value, (int, float)):
        if not 0 <= value < 4:  # reject Excel calendar serials
            return None
        return round(float(value) * 24 * 60)
    if isinstance(value, str):
        match = re.fullmatch(r"\s*(\d{1,2}):([0-5]\d)\s*", value)
        if match:
            hours = int(match.group(1))
            if hours < 48:
                return hours * 60 + int(match.group(2))
    return None


def _text_parts(raw: str) -> tuple[str, str, str] | None:
    if not LIVE_RE.match(raw):
        return None
    title = " ".join(LIVE_RE.sub("", raw, count=1).split()).strip(" .")
    if not title:
        return None
    prefix, dot, rest = title.partition(".")
    sport = SPORTS.get(prefix.strip().casefold())
    if not dot or not sport:
        return None  # an opaque LIVE entry must be reviewed, not invented
    event = {
        "raw_title": title, "title": title,
        "sport": sport, "tournament": rest.strip().split(",")[0].strip(),
    }
    if event_is_editorial_or_replay(event):
        return None
    # Extract participant pair only when it is the last complete comma clause.
    segments = [x.strip() for x in rest.split(",")]
    tail = segments[-1] if segments else rest
    if len(segments) >= 2 and re.search(r"\s+[-–]\s+", tail):
        title_for_event = re.sub(r"\s+[–]\s+", " - ", tail)
        tournament = ", ".join(segments[:-1]).strip()
    else:
        title_for_event = rest.strip()
        tournament = segments[0] if segments else ""
    return sport, tournament, title_for_event


def _read_programs(sheet, channel: str, filename: str, year: int) -> tuple[list[dict], set[str], int]:
    """Accept two audited layouts: Setanta1 A/B, others B/C."""
    starts: list[dict] = []
    seen_dates: set[str] = set()
    current_day: date | None = None
    previous_minutes: int | None = None
    rollover = 0
    programmes = 0
    layout = ("A", "B") if channel == CHANNELS["setanta1"] else ("B", "C")
    t_idx = 0 if layout[0] == "A" else 1
    title_idx = 1 if layout[1] == "B" else 2
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, MAX_SHEET_ROWS),
                               min_col=1, max_col=5, values_only=True):
        title_cell = row[title_idx]
        for candidate in (row[1], row[2]):
            header = _day_header(candidate, year)
            if header:
                current_day = header
                seen_dates.add(header.isoformat())
                previous_minutes = None
                rollover = 0
                break
        if current_day is None or not isinstance(title_cell, str):
            continue
        minutes = _clock_minutes(row[t_idx])
        if minutes is None:
            continue
        # Supplier programmes run from 06:00 to next-day 05:59. The Excel
        # serial sometimes wraps back to 0.20 after a preceding 1.08 row.
        base_minutes = minutes % 1440
        if previous_minutes is not None and minutes < 1440:
            if previous_minutes >= 18 * 60 and base_minutes < 6 * 60:
                rollover = max(rollover, 1)
        day_offset = max(minutes // 1440, rollover)
        absolute_minutes = day_offset * 1440 + base_minutes
        if previous_minutes is not None and absolute_minutes < previous_minutes:
            # An out-of-order spreadsheet row is not trustworthy as a slot.
            continue
        previous_minutes = absolute_minutes
        programmes += 1
        start = datetime.combine(current_day, time.min, tzinfo=KZ) + timedelta(minutes=absolute_minutes)
        starts.append({"start": start, "raw": " ".join(title_cell.split()),
                       "duration": _clock_minutes(row[3]) if title_idx == 2 else None})
    return starts, seen_dates, programmes


def parse_epg_xlsx(data: bytes, filename: str, *, today: date | None = None) -> ParsedEPG:
    if not filename.casefold().endswith(".xlsx"):
        raise InvalidEPG("Поддерживаются только файлы .xlsx")
    if not data or len(data) > MAX_WORKBOOK_BYTES:
        raise InvalidEPG("Размер Excel превышает 6 МБ или файл пуст")
    # The supplied filename is untrusted, even after successful MIME parsing.
    filename = PurePath(filename.replace("\\", "/")).name
    if len(filename) > 180 or not filename.strip():
        raise InvalidEPG("Недопустимое имя вложения")
    channel = detect_channel(filename)
    try:
        with ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
            if (len(infos) > 512
                    or sum(item.file_size for item in infos) > 40 * 1024 * 1024
                    or any(item.file_size > 30 * 1024 * 1024 for item in infos)
                    or any(item.filename.casefold().endswith("vbaproject.bin")
                           for item in infos)):
                raise InvalidEPG("Подозрительное содержимое XLSX")
    except BadZipFile as exc:
        raise InvalidEPG("Повреждённый XLSX") from exc
    today = today or datetime.now(KZ).date()
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise InvalidEPG("Не удалось прочитать XLSX") from exc
    rows = []
    days: set[str] = set()
    count = 0
    try:
        year = _year(filename, today)
        for sheet in workbook.worksheets:
            slot_rows, sheet_days, programmes = _read_programs(sheet, channel, filename, year)
            days.update(sheet_days)
            count += programmes
            rows.extend(slot_rows)
    finally:
        workbook.close()
    if not days or count < 1:
        raise InvalidEPG("В файле не найдены датированные строки программы")
    # Sort without silently merging different programmes; next chronological
    # programme is an EPG end boundary, not proof of an event's actual end.
    rows.sort(key=lambda row: row["start"])
    events: list[dict] = []
    for index, item in enumerate(rows):
        parts = _text_parts(item["raw"])
        if parts is None:
            continue
        sport, tournament, title = parts
        start = item["start"]
        next_start = rows[index + 1]["start"] if index + 1 < len(rows) else None
        duration = item["duration"]
        end = None
        method = None
        if duration is not None and 1 <= duration <= 8 * 60:
            end = start + timedelta(minutes=duration)
            method = "provider_duration"
        elif next_start is not None and 0 < (next_start - start).total_seconds() <= 8 * 3600:
            end = next_start
            method = "next_program"
        record = {
            "source": SOURCE_KEY[channel], "source_url": "",
            "provider_source": "supplier_excel",
            "channel": channel, "date": start.date().isoformat(),
            "time": start.strftime("%H:%M"), "title": title,
            "raw_title": item["raw"], "sport": sport,
            "tournament": tournament, "is_sport_event": True,
            "is_live": True, "is_live_broadcast": True,
            "live_state": "live", "live_evidence_method": "provider_live_text",
            "live_evidence_value": "LIVE. in supplier XLSX",
            "source_timezone": "Asia/Almaty",
            "source_start_at": start.isoformat(),
        }
        if end is not None and method:
            record["estimated_broadcast_end_date"] = end.date().isoformat()
            record["estimated_broadcast_end"] = end.strftime("%H:%M")
            record["end_estimation_method"] = method
        events.append(record)
    # XLSX may repeat the same channel event on multiple sheets; preserve
    # different real start times and real channel identities.
    unique = {
        (e["date"], e["time"], e["channel"], e["title"].casefold()): e
        for e in events
    }
    return ParsedEPG(channel=channel, filename=filename,
                     content_hash=hashlib.sha256(data).hexdigest(),
                     scope_dates=tuple(sorted(days)), all_programmes=count,
                     events=tuple(unique.values()))


def initialize_epg_imports(database: SLPDatabase) -> None:
    with database._connect() as conn:
        conn.executescript("""
            CREATE TABLE IF NOT EXISTS epg_imports (
                file_hash TEXT PRIMARY KEY,
                channel TEXT NOT NULL,
                filename TEXT NOT NULL,
                first_date TEXT NOT NULL,
                last_date TEXT NOT NULL,
                programme_count INTEGER NOT NULL,
                live_count INTEGER NOT NULL,
                imported_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_epg_imports_channel
                ON epg_imports(channel, imported_at DESC);
            CREATE TABLE IF NOT EXISTS epg_import_days (
                file_hash TEXT NOT NULL REFERENCES epg_imports(file_hash),
                scope_date TEXT NOT NULL,
                PRIMARY KEY(file_hash, scope_date)
            );
        """)


def imported_epg_status(database: SLPDatabase, *, first: date, last: date) -> list[dict]:
    """Distinguish unavailable, old, and partially covered Excel sources."""
    initialize_epg_imports(database)
    result = []
    with database._connect() as conn:
        for channel in CHANNELS.values():
            row = conn.execute(
                "SELECT * FROM epg_imports WHERE channel=? "
                "ORDER BY imported_at DESC LIMIT 1", (channel,)
            ).fetchone()
            if row is None:
                result.append({"channel": channel, "status": "missing",
                               "coverage": [], "live_events": 0})
                continue
            days = [x[0] for x in conn.execute(
                "SELECT scope_date FROM epg_import_days WHERE file_hash=? ORDER BY scope_date",
                (row["file_hash"],),
            ).fetchall()]
            covered = [day for day in days if first.isoformat() <= day <= last.isoformat()]
            state = ("ready" if len(covered) == (last-first).days+1 else
                     "partial" if covered else "outdated")
            result.append({"channel": channel, "status": state,
                           "coverage": days, "live_events": row["live_count"],
                           "filename": row["filename"], "imported_at": row["imported_at"]})
    return result


def import_parsed_epg(database: SLPDatabase, parsed: ParsedEPG) -> dict:
    """Idempotent on attachment SHA; preserve other channels and unrelated days."""
    initialize_epg_imports(database)
    with database._connect() as conn:
        exists = conn.execute("SELECT 1 FROM epg_imports WHERE file_hash=?",
                              (parsed.content_hash,)).fetchone()
    if exists:
        return {"status": "already_imported", "channel": parsed.channel,
                "live_events": len(parsed.events), "dates": list(parsed.scope_dates)}
    by_day = defaultdict(list)
    for event in parsed.events:
        by_day[event["date"]].append(event)
    source = SOURCE_KEY[parsed.channel]
    # No destructive sweep for dates absent from this attachment.
    for day in parsed.scope_dates:
        database.upsert_source_snapshot(
            run_id="epg:" + parsed.content_hash[:20],
            source=source, scope_date=day, events=by_day.get(day, []),
        )
    with database._connect() as conn:
        conn.execute(
            "INSERT INTO epg_imports(file_hash,channel,filename,first_date,last_date,"
            "programme_count,live_count,imported_at) VALUES(?,?,?,?,?,?,?,?)",
            (parsed.content_hash, parsed.channel, parsed.filename,
             parsed.scope_dates[0], parsed.scope_dates[-1],
             parsed.all_programmes, len(parsed.events),
             datetime.now(KZ).isoformat()),
        )
        conn.executemany(
            "INSERT INTO epg_import_days(file_hash,scope_date) VALUES(?,?)",
            [(parsed.content_hash, day) for day in parsed.scope_dates],
        )
    return {"status": "imported", "channel": parsed.channel,
            "live_events": len(parsed.events), "dates": list(parsed.scope_dates)}
