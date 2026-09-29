"""Focused regression for the two actual supplier XLSX layouts. No Gmail or network."""
import hashlib
from datetime import date
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from services.epg_excel import (
    InvalidEPG, import_parsed_epg, imported_epg_status, parse_epg_xlsx,
    parse_epg_xlsx_channels, preview_parsed_epg, workbook_fingerprint,
)
from storage.database import SLPDatabase
from cloudrun_web import event_rows


def workbook_bytes(rows):
    workbook = Workbook()
    ws = workbook.active
    for col, row, value in rows:
        ws.cell(row, col, value)
    result = BytesIO()
    workbook.save(result)
    workbook.close()
    return result.getvalue()


class EPGExcelTests(unittest.TestCase):
    def test_structure_fingerprint_ignores_weekly_fixture_text(self):
        first = workbook_bytes([
            (3, 3, "29 сентября"), (2, 4, "AST"),
            (2, 5, .5), (3, 5, "LIVE. Футбол. Лига, А - Б"),
        ])
        second = workbook_bytes([
            (3, 3, "6 октября"), (2, 4, "AST"),
            (2, 5, .6), (3, 5, "LIVE. Теннис. Турнир, В - Г"),
        ])
        self.assertEqual(
            workbook_fingerprint(first, "EPG QSport Arena.xlsx"),
            workbook_fingerprint(second, "сетка Канала.xlsx"),
        )

    def test_confirmed_format_accepts_unlabelled_sheet_but_not_conflict(self):
        unlabelled = workbook_bytes([
            (3, 3, "29 сентября"), (2, 4, "AST"),
            (2, 5, .5), (3, 5, "LIVE. Футбол. Лига, А - Б"),
        ])
        parsed = parse_epg_xlsx_channels(
            unlabelled, "сетка Канала.xlsx", today=date(2026, 9, 29),
            confirmed_channel="Q ARENA",
        )
        self.assertEqual(parsed[0].channel, "Q ARENA")
        conflict = workbook_bytes([
            (2, 1, "SETANTA SPORTS 1 KAZAKHSTAN"),
            (2, 4, "29 сентября"), (1, 5, "AST"),
            (1, 6, .5), (2, 6, "LIVE. Футбол. Лига, А - Б"),
        ])
        with self.assertRaisesRegex(InvalidEPG, "противоречит"):
            parse_epg_xlsx_channels(
                conflict, "сетка Канала.xlsx", today=date(2026, 9, 29),
                confirmed_channel="Q ARENA",
            )

    def test_setanta1_one_sheet_live_only_and_epg_end(self):
        data = workbook_bytes([
            (2, 4, "29 сентября"), (1, 5, "AST"),
            (1, 6, 0.50), (2, 6, "LIVE. Футбол. АПЛ, 6 тур, Команда А - Команда Б"),
            (1, 7, 0.60), (2, 7, "Футбол. Повтор матча"),
            (1, 8, 0.70), (2, 8, "LIVE. Студийная программа"),
        ])
        result = parse_epg_xlsx(data,
            "EPG Setanta Sports 1 Kazakhstan 29.09.26 - 05.10.26_MEDIA.xlsx",
            today=date(2026, 9, 29))
        self.assertEqual(result.channel, "SETANTA SPORTS 1")
        self.assertEqual(len(result.events), 1)
        event = result.events[0]
        self.assertEqual((event["date"], event["time"]), ("2026-09-29", "12:00"))
        self.assertEqual(event["title"], "Команда А - Команда Б")
        self.assertEqual(event["sport"], "Футбол")
        self.assertEqual(event["tournament"], "АПЛ, 6 тур")
        self.assertEqual(event["end_estimation_method"], "next_program")
        self.assertEqual(event["estimated_broadcast_end"], "14:24")

    def test_qsport_multisheet_duration_and_midnight(self):
        workbook = Workbook()
        first = workbook.active
        first["C3"] = "30 сентября"
        first["B4"] = "AST"
        first["B5"] = 0.99
        first["C5"] = "LIVE. Футбол. Лига, 4 тур, Команда А – Команда Б"
        first["D5"] = 0.10
        first["B6"] = 1.09
        first["C6"] = "Обзор матчей"
        later = workbook.create_sheet()
        later["C3"] = "1 октября"
        later["B4"] = "AST"
        later["B5"] = 0.50
        later["C5"] = "LIVE. Фигурное катание. Финал, женщины"
        later["D5"] = 0.08
        data = BytesIO()
        workbook.save(data)
        workbook.close()
        result = parse_epg_xlsx(data.getvalue(),
            "EPG QSport Arena 28.09.26 - 04.10.26_MEDIA.xlsx",
            today=date(2026, 9, 29))
        self.assertEqual(result.channel, "Q ARENA")
        self.assertEqual(len(result.events), 2)
        self.assertEqual(result.events[0]["date"], "2026-09-30")
        self.assertEqual(result.events[0]["time"], "23:46")
        self.assertEqual(result.events[0]["estimated_broadcast_end_date"], "2026-10-01")
        self.assertEqual(result.events[0]["estimated_broadcast_end"], "02:10")
        self.assertEqual(result.events[1]["date"], "2026-10-01")

    def test_year_boundary_across_separate_supplier_sheets(self):
        """A weekly Dec-Jan file must not place January matches in last year."""
        workbook = Workbook()
        december = workbook.active
        december["C3"] = "31 декабря"
        december["B4"] = "AST"
        december["B5"] = .75
        december["C5"] = "LIVE. Футбол. Лига, Команда А - Команда Б"
        january = workbook.create_sheet("Неделя 2")
        january["C3"] = "1 января"
        january["B4"] = "AST"
        january["B5"] = .5
        january["C5"] = "LIVE. Теннис. ATP 250, Финал"
        contents = BytesIO()
        workbook.save(contents)
        workbook.close()
        parsed = parse_epg_xlsx_channels(
            contents.getvalue(),
            "EPG QSport Arena 30.12.26 - 05.01.27_MEDIA.xlsx",
            today=date(2026, 12, 30),
        )[0]
        self.assertEqual(
            [(event["date"], event["time"]) for event in parsed.events],
            [("2026-12-31", "18:00"), ("2027-01-01", "12:00")],
        )
        self.assertIn("2027-01-01", parsed.scope_dates)
        with tempfile.TemporaryDirectory() as folder:
            database = SLPDatabase(Path(folder) / "epg.sqlite")
            result = import_parsed_epg(database, parsed)
            self.assertEqual(result["status"], "imported")
            self.assertEqual(len(database.load_active_source_snapshot(
                "email_epg_qarena", "2027-01-01",
            )), 1)

    def test_year_boundary_without_filename_year_needs_previous_december(self):
        """An unlabelled next-year date is inferred only within one sheet."""
        data = workbook_bytes([
            (3, 3, "31 декабря"), (2, 4, "AST"),
            (2, 5, .5), (3, 5, "LIVE. Футбол. Лига, Команда А - Команда Б"),
            (3, 6, "1 января"), (2, 7, "AST"),
            (2, 8, .5), (3, 8, "LIVE. Теннис. ATP 250, Финал"),
        ])
        parsed = parse_epg_xlsx(
            data, "EPG QSport Arena.xlsx", today=date(2026, 12, 31),
        )
        self.assertEqual(
            [event["date"] for event in parsed.events],
            ["2026-12-31", "2027-01-01"],
        )

    def test_real_setanta_formula_one_and_tennis_title_shapes(self):
        data = workbook_bytes([
            (2, 4, "3 октября"), (1, 5, "AST"),
            (1, 6, 0.50),
            (2, 6, "LIVE. Формула 1: Гран-при Бахрейна - Квалификация"),
            (1, 7, 0.60),
            (2, 7, "LIVE. Теннис. ATP 250 Ханчжоу: Финал"),
            (1, 8, 0.70), (2, 8, "Обзор матчей"),
        ])
        result = parse_epg_xlsx(
            data,
            "EPG Setanta Sports 1 Kazakhstan 29.09.26 - 05.10.26_MEDIA.xlsx",
            today=date(2026, 9, 29),
        )
        self.assertEqual(len(result.events), 2)
        f1, tennis = result.events
        self.assertEqual(
            (f1["sport"], f1["tournament"], f1["title"]),
            ("Автоспорт", "Формула-1. Гран-при Бахрейна", "Квалификация"),
        )
        self.assertEqual(
            (tennis["sport"], tennis["tournament"], tennis["title"]),
            ("Теннис", "ATP 250 Ханчжоу", "Финал"),
        )

    def test_missing_live_is_not_invented_and_hash_idempotent(self):
        data = workbook_bytes([
            (3, 3, "1 октября"), (2, 4, "AST"),
            (2, 5, 0.50), (3, 5, "Футбол. Старый повтор"),
        ])
        result = parse_epg_xlsx(data, "EPG SETANTA QAZAQSTAN 29.09.26 - 05.10.26_MEDIA.xlsx",
                                today=date(2026, 9, 29))
        self.assertEqual(result.channel, "SETANTA SPORTS KZ")
        self.assertEqual(result.events, ())
        with tempfile.TemporaryDirectory() as folder:
            db = SLPDatabase(Path(folder) / "db.sqlite")
            self.assertEqual(import_parsed_epg(db, result)["status"], "imported")
            self.assertEqual(import_parsed_epg(db, result)["status"], "already_imported")
            statuses = imported_epg_status(db, first=date(2026, 10, 1),
                                           last=date(2026, 10, 1))
            kz = next(x for x in statuses if x["channel"] == "SETANTA SPORTS KZ")
            self.assertEqual(kz["status"], "ready")
            self.assertEqual(kz["live_events"], 0)
            self.assertEqual(db.active_event_count(), 0)

    def test_midnight_event_not_lost_when_next_day_header_absent(self):
        data = workbook_bytes([
            (3, 3, "29 сентября"), (2, 4, "AST"),
            (2, 5, 1.05), (3, 5, "LIVE. Футбол. Лига, 5 тур, Команда А - Команда Б"),
            (2, 6, 1.13), (3, 6, "Обзор футбола"),
        ])
        parsed = parse_epg_xlsx(
            data, "EPG QSport Arena 28.09.26 - 04.10.26_MEDIA.xlsx",
            today=date(2026, 9, 29),
        )
        self.assertEqual(parsed.events[0]["date"], "2026-09-30")
        self.assertIn("2026-09-30", parsed.scope_dates)
        with tempfile.TemporaryDirectory() as folder:
            db = SLPDatabase(Path(folder) / "db.sqlite")
            import_parsed_epg(db, parsed)
            self.assertEqual(len(db.load_active_source_snapshot(
                "email_epg_qarena", "2026-09-30"
            )), 1)

    def test_preview_shows_changed_time_without_inventing_cancellation(self):
        def supplier(minutes):
            return workbook_bytes([
                (3, 3, "29 сентября"), (2, 4, "AST"),
                (2, 5, minutes),
                (3, 5, "LIVE. Футбол. АПЛ, Команда А - Команда Б"),
                (2, 6, minutes + .15),
                (3, 6, "Новости"),
            ])
        with tempfile.TemporaryDirectory() as folder:
            db = SLPDatabase(Path(folder) / "db.sqlite")
            initial = parse_epg_xlsx(
                supplier(.5), "EPG QSport Arena 28.09.26 - 04.10.26_MEDIA.xlsx",
                today=date(2026, 9, 29),
            )
            import_parsed_epg(db, initial)
            changed = parse_epg_xlsx(
                supplier(.55), "EPG QSport Arena 28.09.26 - 04.10.26_MEDIA.xlsx",
                today=date(2026, 9, 29),
            )
            preview = preview_parsed_epg(db, changed)
            self.assertEqual(preview["counts"]["time_changed"], 1)
            self.assertEqual(preview["changes"][0]["before_times"], ["12:00"])
            self.assertEqual(preview["changes"][0]["after_times"], ["13:12"])
            self.assertEqual(
                len(db.load_active_source_snapshot("email_epg_qarena", "2026-09-29")),
                1,
            )

    def test_multisheet_workbook_routes_stations_separately(self):
        from services.epg_excel import parse_epg_xlsx_channels
        wb = Workbook()
        first = wb.active
        first.title = "Q ARENA"
        first["C3"] = "29 сентября"
        first["B4"] = "AST"
        first["B5"] = .5
        first["C5"] = "LIVE. Футбол. КПЛ, Команда А - Команда Б"
        first["B6"] = .6
        first["C6"] = "Новости"
        second = wb.create_sheet("Q LEAGUE")
        second["C3"] = "29 сентября"
        second["B4"] = "AST"
        second["B5"] = .55
        second["C5"] = "LIVE. Футбол. КПЛ, Команда В - Команда Г"
        second["B6"] = .65
        second["C6"] = "Новости"
        stream = BytesIO()
        wb.save(stream)
        wb.close()
        content = stream.getvalue()
        with self.assertRaises(InvalidEPG):
            parse_epg_xlsx(content, "сетка Канала.xlsx",
                           today=date(2026, 9, 29))
        results = parse_epg_xlsx_channels(
            content, "сетка Канала.xlsx",
            today=date(2026, 9, 29), context="Сетки QSport",
        )
        self.assertEqual([r.channel for r in results], ["Q ARENA", "Q LEAGUE"])
        self.assertEqual([len(r.events) for r in results], [1, 1])
        self.assertNotEqual(results[0].content_hash, results[1].content_hash)
        with tempfile.TemporaryDirectory() as folder:
            db = SLPDatabase(Path(folder) / "slp.db")
            for item in results:
                import_parsed_epg(db, item)
            self.assertEqual(db.active_event_count(), 2)

    def test_multistation_unlabelled_sheet_is_not_guessed(self):
        from services.epg_excel import parse_epg_xlsx_channels
        wb = Workbook()
        first = wb.active
        first.title = "Q ARENA"
        second = wb.create_sheet("Q LEAGUE")
        wb.create_sheet("Неделя")
        output = BytesIO()
        wb.save(output)
        wb.close()
        with self.assertRaisesRegex(InvalidEPG, "не каждый лист"):
            parse_epg_xlsx_channels(
                output.getvalue(), "сетка Канала.xlsx",
                today=date(2026, 9, 29),
            )

    def test_generic_filename_uses_workbook_identity(self):
        data = workbook_bytes([
            (2, 1, "ПРОГРАММА ПЕРЕДАЧ ТЕЛЕКАНАЛА SETANTA SPORTS 1 KAZAKHSTAN"),
            (2, 4, "29 сентября"), (1, 5, "AST"),
            (1, 6, 0.50),
            (2, 6, "LIVE. Футбол. АПЛ, 6 тур, Команда А - Команда Б"),
        ])
        result = parse_epg_xlsx(
            data, "сетка Канала.xlsx", today=date(2026, 9, 29),
            context="Sabina Nazarova Setanta Sports",
        )
        self.assertEqual(result.channel, "SETANTA SPORTS 1")
        self.assertEqual(len(result.events), 1)

    def test_workbook_identity_beats_ambiguous_email_context(self):
        data = workbook_bytes([
            (3, 1, "Q ARENA"),
            (3, 3, "30 сентября"), (2, 4, "AST"),
            (2, 5, 0.50),
            (3, 5, "LIVE. Футбол. КПЛ, 27 тур, Команда А - Команда Б"),
        ])
        result = parse_epg_xlsx(
            data, "сетка Канала.xlsx", today=date(2026, 9, 29),
            context="Расписания Q LEAGUE, Q ARENA, Q FOOTBALL",
        )
        self.assertEqual(result.channel, "Q ARENA")

    def test_unknown_channel_rejected(self):
        with self.assertRaises(InvalidEPG):
            parse_epg_xlsx(b"not a spreadsheet", "mystery.xlsx")

    def test_unlabelled_regional_sheet_is_not_mislabeled_as_known_channel(self):
        book = Workbook()
        known = book.active
        known.title = "Q ARENA"
        known["C3"] = "30 сентября"
        known["B4"] = "AST"
        known["B5"] = 0.50
        known["C5"] = "LIVE. Футбол. КПЛ, 27 тур, Команда А - Команда Б"
        other = book.create_sheet("Другой регион")
        other["C3"] = "30 сентября"
        other["B4"] = "AST"
        other["B5"] = 0.60
        other["C5"] = "LIVE. Футбол. КПЛ, 27 тур, Команда В - Команда Г"
        payload = BytesIO()
        book.save(payload)
        book.close()
        with self.assertRaisesRegex(InvalidEPG, "без указания канала"):
            parse_epg_xlsx_channels(
                payload.getvalue(), "сетка Канала.xlsx",
                today=date(2026, 9, 29),
                context="Расписание Q ARENA",
            )

    def test_legacy_xls_routes_by_actual_workbook_and_retains_original_hash(self):
        # A converted workbook is injected here; the real BIFF decoder is
        # exercised by the dependency at integration time with supplier files.
        xlsx = workbook_bytes([
            (3, 1, "Q ARENA"), (3, 3, "30 сентября"),
            (2, 4, "AST"), (2, 5, 0.5),
            (3, 5, "LIVE. Футбол. КПЛ, 27 тур, Команда А - Команда Б"),
        ])
        legacy = b"sample-legacy-biff-bytes"
        with patch("services.epg_excel._legacy_xls_to_xlsx", return_value=xlsx):
            result = parse_epg_xlsx_channels(
                legacy, "сетка Канала.xls", today=date(2026, 9, 29),
                context="Письмо Q LEAGUE и Q ARENA",
            )
            direct = parse_epg_xlsx(
                legacy, "сетка Канала.xls", today=date(2026, 9, 29),
            )
        self.assertEqual(len(result), 1)
        self.assertEqual(result[0].channel, "Q ARENA")
        self.assertEqual(result[0].filename, "сетка Канала.xls")
        original_hash = hashlib.sha256(b"Q ARENA\0" + legacy).hexdigest()
        self.assertEqual(result[0].content_hash, original_hash)
        self.assertEqual(direct.content_hash, original_hash)
        self.assertEqual(result[0].events[0]["time"], "12:00")

    def test_partial_supplier_update_preserves_coverage_from_other_files(self):
        with tempfile.TemporaryDirectory() as folder:
            db = SLPDatabase(Path(folder) / "db.sqlite")
            filename = "EPG Setanta Sports 2 Kazakhstan 29.09.26 - 05.10.26_MEDIA.xlsx"
            for day, title in [
                ("30 сентября", "Команда А - Команда Б"),
                ("1 октября", "Команда В - Команда Г"),
            ]:
                grid = workbook_bytes([
                    (3, 3, day), (2, 4, "AST"), (2, 5, 0.5),
                    (3, 5, "LIVE. Футбол. АПЛ, 6 тур, " + title),
                    (2, 6, 0.6), (3, 6, "Обзор матча"),
                ])
                parsed = parse_epg_xlsx(grid, filename, today=date(2026, 9, 29))
                self.assertEqual(import_parsed_epg(db, parsed)["status"], "imported")
            status = next(item for item in imported_epg_status(
                db, first=date(2026, 9, 30), last=date(2026, 10, 1)
            ) if item["channel"] == "SETANTA SPORTS 2")
            self.assertEqual(status["status"], "ready")
            self.assertIn("2026-09-30", status["coverage"])
            self.assertIn("2026-10-01", status["coverage"])
            self.assertEqual(status["live_events"], 2)

    def test_unreadable_legacy_xls_does_not_create_live(self):
        with self.assertRaises(InvalidEPG):
            parse_epg_xlsx_channels(b"not-a-real-biff-workbook", "schedule.xls")

    def test_simulcast_only_different_channels_matching_event(self):
        with tempfile.TemporaryDirectory() as folder:
            db = SLPDatabase(Path(folder) / "db.sqlite")
            def row(channel, tm, tournament="Лига", name="Команда А - Команда Б"):
                return {
                    "source": "unit", "date": "2026-10-03", "time": tm,
                    "channel": channel, "sport": "Футбол", "tournament": tournament,
                    "title": name, "raw_title": name,
                    "is_live_broadcast": True, "is_sport_event": True,
                }
            db.upsert_source_snapshot(
                run_id="t", source="x", scope_date="2026-10-03",
                events=[
                    row("EUROSPORT 1", "18:00"),
                    row("EUROSPORT 2", "18:10"),
                    row("EUROSPORT 1", "19:00"),
                    row("EUROSPORT 2", "20:00", "Кубок"),
                ],
            )
            events = event_rows(db, "2026-10-03", "2026-10-03")
            self.assertEqual(len(events), 3)
            first = events[0]
            self.assertEqual(first["channels"], ["EUROSPORT 1", "EUROSPORT 2"])
            self.assertEqual(len(first["broadcasts"]), 2)
            self.assertEqual(events[1]["channels"], ["EUROSPORT 1"])
            self.assertEqual(events[2]["channels"], ["EUROSPORT 2"])


if __name__ == "__main__":
    unittest.main()
