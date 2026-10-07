"""Focused web-only regression; does not contact external sites or send mail."""
import asyncio
import base64
import hashlib
import json
import os
import tempfile
import unittest
from pathlib import Path
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from fastapi import HTTPException
from openpyxl import Workbook
from starlette.datastructures import UploadFile

from services.editorial_export import HEADERS

import cloudrun_web as web
from storage.database import SLPDatabase


class WebTests(unittest.TestCase):
    def test_session_and_named_avatar(self):
        original = web.SESSION_SECRET
        web.SESSION_SECRET = "a-private-signing-secret-at-least-32-bytes"
        old = os.environ.get("SPORT_WEB_USERS")
        try:
            pw = "new-long-private-password"
            salt = b"1234567890123456"
            hashed = "scrypt$" + base64.b64encode(salt).decode() + "$" + base64.b64encode(
                hashlib.scrypt(pw.encode(), salt=salt, n=16384, r=8, p=1, dklen=32)
            ).decode()
            os.environ["SPORT_WEB_USERS"] = json.dumps({
                "Дания": {"password_hash": hashed, "role": "editor", "avatar": "🇩🇰"}
            }, ensure_ascii=False)
            self.assertTrue(web._password_matches(pw, hashed))
            fingerprint = hashlib.sha256(hashed.encode()).hexdigest()
            request = SimpleNamespace(cookies={
                "sport_web_session": web._session("Дания", fingerprint)
            })
            self.assertEqual(web.current_user(request)["avatar"], "🇩🇰")
            self.assertFalse(web._password_matches("wrong", hashed))
        finally:
            web.SESSION_SECRET = original
            if old is None:
                os.environ.pop("SPORT_WEB_USERS", None)
            else:
                os.environ["SPORT_WEB_USERS"] = old

    def test_supplied_profile_avatars_and_app_icon_assets(self):
        # The authenticated identity selects an application-owned image.
        self.assertEqual(web.profile_avatar("Skoomaholic"),
                         "/assets/skoomaholic.webp")
        self.assertEqual(web.profile_avatar("Дания"),
                         "/assets/daniya.webp")
        self.assertEqual(web.profile_avatar("Вадим"),
                         "/assets/vadim.webp")
        self.assertEqual(web.profile_avatar("USER4"), "")
        assets = Path(__file__).resolve().parents[1] / "cloudrun_ui" / "assets"
        expected = {
            "app-icon.png": b"\x89PNG",
            "skoomaholic.webp": b"RIFF",
            "daniya.webp": b"RIFF",
            "vadim.webp": b"RIFF",
        }
        hashes = []
        for name, prefix in expected.items():
            raw = (assets / name).read_bytes()
            self.assertTrue(raw.startswith(prefix), name)
            self.assertGreater(len(raw), 500, name)
            hashes.append(hashlib.sha256(raw).digest())
        self.assertEqual(len(hashes), len(set(hashes)))

    def test_ui_names_channels_without_swapping_images(self):
        import re
        index = (Path(__file__).resolve().parents[1] /
                 "cloudrun_ui" / "index.html").read_text()
        logos_js = (Path(__file__).resolve().parents[1] /
                    "cloudrun_ui" / "logos.js").read_text()
        web_source = (Path(__file__).resolve().parents[1] /
                      "cloudrun_web.py").read_text()
        self.assertIn('href="/assets/app-icon.png"', index)
        self.assertIn('function channelLogo(channel)', index)
        self.assertNotIn('className="logo-variant"', index)
        self.assertNotIn('channel==="KHL PRIME"?"PRIME":"HD"', index)
        logo_map = re.search(
            r'Object.freeze\((\{.*\})\);', logos_js, re.DOTALL)
        self.assertIsNotNone(logo_map)
        logos = json.loads(logo_map.group(1))
        self.assertNotEqual(logos["KHL PRIME"], logos["KHL HD"])
        self.assertNotEqual(logos["SETANTA SPORTS 1"], logos["SETANTA SPORTS 2"])
        self.assertNotEqual(logos["EUROSPORT 1"], logos["EUROSPORT 2"])
        self.assertIn(".day.today", index)
        self.assertIn('const today=kzToday();', index)
        self.assertIn('date:""', index)
        self.assertIn('id="allDates">Все даты<', index)
        self.assertIn("COLLECTION_QUOTES", index)
        self.assertIn('setInterval(()=>{', index)
        self.assertIn('},12000);', index)
        self.assertIn("collection-quote", index)
        self.assertIn("text-align:center", index)
        self.assertIn('authorLine.textContent="(с) "+author', index)
        self.assertIn('replace(/\\.+$/,"")', index)
        self.assertIn('api("/api/collect/status")', index)
        self.assertIn('await api("/api/collect","POST"', index)
        self.assertIn("status.result", index)
        self.assertIn("Наврал - ещё собираю", web_source)
        self.assertIn("Вот-вот", web_source)
        self.assertIn("Расписание собрано:", web_source)
        self.assertIn("Ошибка #", web_source)
        self.assertIn("const combined=new Map()", index)
        self.assertNotIn('[...(data.websites||[]),...(data.excel||[])]', index)
        self.assertIn('approve.textContent="Добавить в расписание"', index)

    def test_source_cards_are_compact_and_collection_retries_429(self):
        index = (
            Path(__file__).resolve().parents[1] /
            "cloudrun_ui" / "index.html"
        ).read_text()
        self.assertIn("function sourceStamp(value)", index)
        # Excel and parsing are two separate indicators; each is "on" only
        # when it has events in the current 7-day window.
        self.assertIn('chip("Excel","on"', index)
        self.assertIn('chip("Парсинг",web.error_reason?"warn":"on"', index)
        self.assertIn('"файл устарел"', index)
        self.assertIn('"нет LIVE-событий"', index)
        self.assertNotIn('origins.join(" + ")', index)
        self.assertNotIn('webState.textContent="Парсинг: "', index)
        self.assertNotIn('mailState.textContent="Почта/Excel: "', index)
        self.assertNotIn("Gmail подключён. Новые EPG попадают", index)
        self.assertIn("if(e.status!==429)throw e", index)
        self.assertIn("attempt<8", index)
        self.assertIn("if(e.status===429)continue", index)
        self.assertIn("setTimeout(resolve,3000)", index)

    def test_idle_status_and_event_cards_are_contextual(self):
        index = (
            Path(__file__).resolve().parents[1] /
            "cloudrun_ui" / "index.html"
        ).read_text()
        web_source = (
            Path(__file__).resolve().parents[1] /
            "cloudrun_web.py"
        ).read_text()
        self.assertIn("function eventMeta(ev)", index)
        self.assertIn('label.textContent=eventMeta(ev)', index)
        self.assertNotIn("Сбор ещё не запускался", web_source)
        self.assertIn("Расписание готово к обновлению", web_source)
        self.assertIn("Событий в базе:", web_source)

    def test_source_inventory_covers_exact_14_channels_without_duplicate_web_cards(self):
        with tempfile.TemporaryDirectory() as directory:
            database = SLPDatabase(Path(directory) / "sport.db")
            request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(
                database=database, backup=None
            )))
            status = web._source_status(request)
            channels = [item["channel"] for item in (
                status["websites"] + status["excel"]
            )]
            self.assertEqual(len(status["websites"]), 14)
            self.assertEqual(len(status["excel"]), 6)
            self.assertEqual(len(set(channels)), 14)
            self.assertEqual(len({item["channel"] for item in status["websites"]}), 14)
            self.assertNotIn("TVGuide (проверенные LIVE)", channels)
            for not_approved in ("FIGHT CLUB", "QAZAQSTAN", "MMA-TV.COM", "БОКС ТВ"):
                self.assertNotIn(not_approved, channels)
            self.assertIn("KHL PRIME", channels)
            self.assertIn("KHL HD", channels)
            self.assertIn("EUROSPORT 1", channels)
            self.assertIn("EUROSPORT 2", channels)

    def test_supplier_mail_does_not_fake_website_health_and_error_is_explained(self):
        with tempfile.TemporaryDirectory() as directory:
            database = SLPDatabase(Path(directory) / "sport.db")
            database.upsert_source_snapshot(
                run_id="mail-only",
                source="email_epg_setanta1",
                scope_date="2026-10-02",
                events=[{
                    "source": "email_epg_setanta1",
                    "source_url": "",
                    "date": "2026-10-02",
                    "time": "20:00",
                    "channel": "SETANTA SPORTS 1",
                    "sport": "Футбол",
                    "tournament": "Премьер-лига",
                    "title": "Арсенал - Челси",
                    "raw_title": "Арсенал - Челси",
                    "is_live": True,
                    "is_live_broadcast": True,
                }],
            )
            request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(
                database=database, backup=None, free_mail_bridge=False
            )))
            with patch.object(web, "tvplus_channel_diagnostics", return_value={
                "Setanta Sports 1": {
                    "loaded_events": 0,
                    "errors": ["API провайдера не ответил: TimeoutError"],
                },
            }):
                status = web._source_status(request)
            item = next(
                entry for entry in status["websites"]
                if entry["channel"] == "SETANTA SPORTS 1"
            )
            self.assertEqual(item["event_count"], 0)
            self.assertEqual(item["status"], "error")
            self.assertIn("тайм-аут соединения", item["error_reason"])

    def test_collect_is_background_and_not_blocked_by_pending_gmail_notice(self):
        with tempfile.TemporaryDirectory() as directory:
            database = SLPDatabase(Path(directory) / "sport.db")
            refresh = AsyncMock(return_value=SimpleNamespace())
            state = SimpleNamespace(
                database=database,
                backup=None,
                collect_lock=asyncio.Lock(),
                collect_task=None,
                collect_progress={},
                free_mail_bridge=False,
                schedule=SimpleNamespace(refresh=refresh),
            )
            request = SimpleNamespace(app=SimpleNamespace(state=state))

            async def scenario():
                response = await web.collect(
                    request,
                    web.CollectOptions(allow_partial=False),
                )
                self.assertTrue(response["accepted"])
                self.assertFalse(response["already_running"])
                self.assertIsNotNone(state.collect_task)
                await state.collect_task
                return state.collect_progress["result"]

            with patch.object(web, "current_user", return_value={"role": "editor"}), \
                 patch.object(web, "origin_guard"), \
                 patch.object(web.gmail, "status", return_value={"connected": False}), \
                 patch.object(web.gmail, "list_notices", return_value=[{
                     "status": "pending",
                     "detected_channel": "SPORT+ Qazaqstan",
                 }]), \
                 patch.object(
                     web, "refresh_vsetv_web_sources",
                     new=AsyncMock(return_value={"stats": []}),
                 ):
                result = asyncio.run(scenario())
            self.assertTrue(result["ok"])
            self.assertEqual(
                result["pending_mail_channels"],
                ["SPORT+ Qazaqstan"],
            )
            refresh.assert_awaited_once()

    def test_collect_without_editor_rights_scans_mail_last_and_never_imports(self):
        with tempfile.TemporaryDirectory() as directory:
            database = SLPDatabase(Path(directory) / "sport.db")
            refresh = AsyncMock(return_value=SimpleNamespace(run_id="run-1"))
            mail_result = {
                "new_attachments": 2,
                "requires_review": 2,
                "auto_imported": 0,
            }
            state = SimpleNamespace(
                database=database,
                backup=None,
                collect_lock=asyncio.Lock(),
                collect_progress={},
                free_mail_bridge=False,
                schedule=SimpleNamespace(refresh=refresh),
            )
            with patch.object(web.gmail, "status", return_value={"connected": True}), \
                 patch.object(web.gmail, "sync_inbox", return_value=mail_result) as sync, \
                 patch.object(web.gmail, "list_notices", return_value=[]), \
                 patch.object(
                     web, "refresh_vsetv_web_sources",
                     new=AsyncMock(return_value={"stats": []}),
                 ):
                result = asyncio.run(
                    web._run_collection(SimpleNamespace(state=state))
                )
            self.assertTrue(result["ok"])
            self.assertEqual(result["mail"]["new_attachments"], 2)
            sync.assert_called_once()
            self.assertFalse(sync.call_args.kwargs["allow_auto_import"])
            refresh.assert_awaited_once()

    def test_same_qazsport_slot_is_canonicalized_and_deduplicated(self):
        with tempfile.TemporaryDirectory() as folder:
            database = SLPDatabase(Path(folder) / "sports.db")

            def event(title, tournament, source):
                return {
                    "source": source,
                    "source_url": "https://example.test/program",
                    "date": "2026-10-01",
                    "time": "23:35",
                    "channel": "QAZSPORT",
                    "sport": "Футбол",
                    "tournament": tournament,
                    "title": title,
                    "raw_title": title,
                    "is_live": True,
                    "is_live_broadcast": True,
                }

            database.upsert_source_snapshot(
                run_id="qazsport-direct",
                source="qazsport",
                scope_date="2026-10-01",
                events=[event(
                    "Дания - Португалия",
                    "Лига наций УЕФА",
                    "qazsport",
                )],
            )
            database.upsert_source_snapshot(
                run_id="tvguide-copy",
                source="tvguide",
                scope_date="2026-10-01",
                events=[event(
                    "Лига наций УЕФА Дания - Португалия",
                    "Лига наций УЕФА Дания - Португалия",
                    "tvguide",
                )],
            )

            rows = web.event_rows(database, "2026-10-01", "2026-10-01")
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["title"], "Дания - Португалия")
            self.assertEqual(rows[0]["tournament"], "Лига наций УЕФА")
            self.assertEqual(rows[0]["sport"], "Футбол")
            self.assertEqual(rows[0]["channel"], "QAZSPORT HD")
            self.assertEqual(rows[0]["broadcasts"][0]["source"], "qazsport")

    def test_supplier_mail_overrides_matching_scrape_and_supplements_missing_event(self):
        with tempfile.TemporaryDirectory() as folder:
            database = SLPDatabase(Path(folder) / "sports.db")

            def event(title, channel, time, source):
                return {
                    "source": source,
                    "source_url": "https://example.test/program",
                    "date": "2026-10-02",
                    "time": time,
                    "channel": channel,
                    "sport": "Футбол",
                    "tournament": "Премьер-лига",
                    "title": title,
                    "raw_title": title,
                    "is_live": True,
                    "is_live_broadcast": True,
                }

            database.upsert_source_snapshot(
                run_id="web-run",
                source="tvguide",
                scope_date="2026-10-02",
                events=[
                    event(
                        "Арсенал - Челси",
                        "SETANTA SPORTS 1",
                        "20:00",
                        "tvguide",
                    ),
                    event(
                        "Ливерпуль - Эвертон",
                        "SETANTA SPORTS 1",
                        "22:00",
                        "tvguide",
                    ),
                ],
            )
            database.upsert_source_snapshot(
                run_id="mail-approved",
                source="email_epg_setanta1",
                scope_date="2026-10-02",
                events=[
                    event(
                        "Арсенал - Челси",
                        "SETANTA SPORTS 1",
                        "20:30",
                        "email_epg_setanta1",
                    ),
                    event(
                        "Брентфорд - Фулхэм",
                        "SETANTA SPORTS 1",
                        "23:00",
                        "email_epg_setanta1",
                    ),
                ],
            )

            rows = web.event_rows(database, "2026-10-02", "2026-10-02")
            by_title = {row["title"]: row for row in rows}

            self.assertEqual(
                by_title["Арсенал - Челси"]["time"],
                "20:30",
            )
            self.assertEqual(
                by_title["Арсенал - Челси"]["broadcasts"][0]["source"],
                "email_epg_setanta1",
            )
            self.assertIn("Брентфорд - Фулхэм", by_title)
            self.assertIn("Ливерпуль - Эвертон", by_title)
            self.assertEqual(len(rows), 3)

    def test_supplier_preview_does_not_import_or_save(self):
        request = SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(
            database=object(), backup=None, collect_lock=asyncio.Lock(),
        )))
        upload = UploadFile(filename="provider.xlsx", file=BytesIO(b"example"))
        parsed = SimpleNamespace(channel="QAZSPORT HD", events=({"title": "LIVE"},))
        summary = {"channel": "QAZSPORT HD", "new_count": 1,
                   "counts": {"new": 1}, "changes": []}
        with patch.object(web, "current_user", return_value={"role": "editor"}), \
             patch.object(web, "origin_guard"), \
             patch.object(web, "parse_supported_epg_channels",
                          return_value=(parsed,)), \
             patch.object(web, "preview_parsed_epg", return_value=summary), \
             patch.object(web, "import_parsed_epg") as importer:
            result = asyncio.run(web.preview_epg(request, upload))
        self.assertEqual(result["previews"], [summary])
        self.assertEqual(result["total_live"], 1)
        self.assertFalse(result["durable_storage"])
        self.assertTrue(result["warning"])
        importer.assert_not_called()

    def test_local_template_upload_without_cloud_storage(self):
        """An admin can configure a template before choosing the final host."""
        workbook = Workbook()
        for col, header in enumerate(HEADERS, 1):
            workbook.active.cell(1, col, header)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "templates" / "approved.xlsx"
            request = SimpleNamespace(
                app=SimpleNamespace(state=SimpleNamespace(
                    backup=None, collect_lock=asyncio.Lock(),
                )),
            )
            upload = UploadFile(
                filename="approved.xlsx", file=BytesIO(output.getvalue()),
            )
            with patch.dict(os.environ, {"SPORT_TEMPLATE_PATH": str(target)}), \
                 patch.object(web, "require_admin", return_value={"role": "admin"}), \
                 patch.object(web, "origin_guard"):
                result = asyncio.run(web.upload_template(request, upload))
                self.assertTrue(result["uploaded"])
                self.assertEqual(result["destination"], "local")
                self.assertIn("постоянного тома", result["warning"])
                self.assertTrue(target.is_file())
                sheet = web.load_template()
                try:
                    self.assertEqual(sheet.active.cell(1, 25).value, HEADERS[24])
                finally:
                    sheet.close()

    def test_local_template_requires_absolute_xlsx_path(self):
        with patch.dict(os.environ, {"SPORT_TEMPLATE_PATH": "templates/sport.xlsx"}):
            with self.assertRaises(HTTPException) as failure:
                web._local_template_path()
        self.assertEqual(failure.exception.status_code, 503)

    def test_scheduler_job_requires_persistence_and_private_oidc(self):
        request = SimpleNamespace(
            app=SimpleNamespace(state=SimpleNamespace(backup=None)),
            headers={},
        )
        with self.assertRaises(HTTPException) as failure:
            asyncio.run(web.scheduled_refresh(request))
        self.assertEqual(failure.exception.status_code, 503)

        request.app.state.backup = object()
        with patch.dict(os.environ, {
            "SPORT_SCHEDULER_SERVICE_ACCOUNT": "cron@example.test"
        }), patch.object(web, "PUBLIC_URL", "https://sport.example"):
            with self.assertRaises(HTTPException) as failure:
                asyncio.run(web.scheduled_refresh(request))
        self.assertEqual(failure.exception.status_code, 401)

    def test_direct_only_dedup_barys_and_time(self):
        with tempfile.TemporaryDirectory() as folder:
            database = SLPDatabase(Path(folder) / "sports.db")
            def event(title, channel):
                return {
                    "source": "qazsport", "source_url": "https://example.test/program",
                    "date": "2026-09-30", "time": "12:00",
                    "channel": channel, "sport": "Футбол",
                    "tournament": "Лига", "title": title,
                    "raw_title": title, "is_live": True,
                    "is_live_broadcast": True,
                }
            database.upsert_source_snapshot(
                run_id="regression", source="qazsport", scope_date="2026-09-30",
                events=[event("Команда А - Команда Б", "QAZSPORT"),
                        event("Барыс - Спартак", "QAZSPORT")],
            )
            database.upsert_source_snapshot(
                run_id="regression", source="tvguide", scope_date="2026-09-30",
                events=[event("Команда А - Команда Б", "SETANTA 1")],
            )
            rows = web.event_rows(database, "2026-09-30", "2026-09-30")
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["channels"], ["QAZSPORT HD", "SETANTA SPORTS 1"])
            self.assertEqual(rows[0]["platform_start_at"][11:16], "11:50")
            self.assertEqual(rows[0]["end_at"][11:16], "14:40")
            self.assertEqual(rows[0]["channel"], "QAZSPORT HD")

    def test_non_sport_and_betting_cleaner(self):
        self.assertEqual(web.clean("ФОНБЕТ Чемпионат КХЛ"), "Чемпионат КХЛ")
        self.assertEqual(web.channel_name("SETANTA 2"), "SETANTA SPORTS 2")


if __name__ == "__main__":
    unittest.main()
