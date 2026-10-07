"""Export must stay downloadable when the Drive/Apps Script bridge is down."""
from datetime import datetime, timedelta
from io import BytesIO
import base64
import hashlib
import unittest
from unittest.mock import patch
from urllib.error import HTTPError

from openpyxl import Workbook, load_workbook

from services.editorial_export import HEADERS
from services import free_apps_script as freebridge
import cloudrun_web as web


def _event(day, clock, title):
    start = datetime.fromisoformat(day + "T" + clock + ":00")
    return {
        "date": day, "time": clock, "title": title,
        "sport": "Футбол", "tournament": "Ла Лига",
        "channel": "SETANTA SPORTS 1",
        "start_at": start.isoformat() + "+05:00",
        "platform_start_at": (start - timedelta(minutes=10)).isoformat() + "+05:00",
        "end_at": (start + timedelta(minutes=150)).isoformat() + "+05:00",
    }


def _approved_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(HEADERS))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class _DownClient:
    def template(self):
        raise freebridge.FreeDriveError(freebridge._http_error_text(404))


class _UpClient:
    def template(self):
        raw = _approved_bytes()
        return {"ok": True, "exists": True,
                "data": base64.b64encode(raw).decode(),
                "sha256": hashlib.sha256(raw).hexdigest()}


class ExportResilienceTests(unittest.TestCase):
    def setUp(self):
        web._TEMPLATE_CACHE.clear()
        self.events = [
            _event("2026-10-07", "21:00", "Команда А - Команда Б"),
            _event("2026-10-06", "19:30", "Команда В - Команда Г"),
        ]

    def tearDown(self):
        web._TEMPLATE_CACHE.clear()

    def _patches(self, client):
        return (
            patch.object(web, "_local_template_path", return_value=None),
            patch.object(web, "GCS_BUCKET", ""),
            patch.object(freebridge, "enabled", return_value=True),
            patch.object(freebridge, "ScriptClient", client),
        )

    def _run(self, client):
        a, b, c, d = self._patches(client)
        with a, b, c, d:
            return web.xlsx_export(self.events)

    def test_bridge_down_without_any_copy_still_exports_all_events(self):
        raw, origin, reason = self._run(_DownClient)
        self.assertEqual(origin, "builtin")
        self.assertIn("HTTP 404", reason)
        sheet = load_workbook(BytesIO(raw)).worksheets[0]
        self.assertEqual(
            tuple(sheet.cell(1, col).value for col in range(1, 26)), HEADERS
        )
        titles = [sheet.cell(row, 6).value for row in range(2, sheet.max_row + 1)]
        self.assertEqual(len([t for t in titles if t]), 2)

    def test_bridge_down_after_one_good_read_uses_cached_approved_copy(self):
        _, origin, _ = self._run(_UpClient)
        self.assertEqual(origin, "approved")
        raw, origin, reason = self._run(_DownClient)
        self.assertEqual(origin, "cached")
        self.assertEqual(reason, "")
        self.assertTrue(raw.startswith(b"PK"))

    def test_http_errors_from_the_bridge_name_the_cause(self):
        self.assertIn("Доступ: Все", freebridge._http_error_text(403))
        self.assertIn("/exec", freebridge._http_error_text(404))
        self.assertIn("HTTP 500", freebridge._http_error_text(500))
        for code in (401, 403, 404, 429, 500, 418):
            self.assertLessEqual(len(freebridge._http_error_text(code)), 150)

    def test_client_reports_status_code_not_bare_httperror(self):
        def boom(*args, **kwargs):
            raise HTTPError("https://script.google.com/x", 403, "Forbidden", {}, None)
        with patch.dict("os.environ", {
            "SPORT_FREE_SCRIPT_URL": "https://script.google.com/macros/s/AAA/exec",
            "SPORT_FREE_SCRIPT_KEY": "k" * 64,
        }), patch.object(freebridge, "urlopen", boom):
            with self.assertRaises(freebridge.FreeDriveError) as caught:
                freebridge.ScriptClient().template()
        self.assertIn("HTTP 403", str(caught.exception))


if __name__ == "__main__":
    unittest.main()
