"""25-column template integrity checks without contacting external services."""
from datetime import datetime
from io import BytesIO
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from services.editorial_export import HEADERS, InvalidTemplate, build_working_xlsx


class EditorialExportTests(unittest.TestCase):
    def _source(self):
        workbook = Workbook()
        sheet = workbook.active
        for col, header in enumerate(HEADERS, 1):
            sheet.cell(1, col).value = header
        source = [
            60000, "24.09", "23:45", "Футбол", "Лига наций УЕФА",
            "Португалия - Уэльс", "Freedom Media 8592",
            datetime(2026, 9, 24, 23, 35), datetime(2026, 9, 25, 1, 45),
            "APPROVED_LIVE_RU", "APPROVED_LIVE_KZ",
            "APPROVED_SOON_RU", "APPROVED_SOON_KZ",
            "", "", "", "Футбол Лига наций УЕФА",
            "APPROVED_ARCH_RU", "APPROVED_ARCH_KZ",
            "ПОРТУГАЛИЯ", "ПОРТУГАЛИЯ", "УЭЛЬС", "УЭЛЬС",
            "Футбол. Лига наций УЕФА", "Футбол. УЕФА Ұлттар лигасы",
        ]
        for col, value in enumerate(source, 1):
            sheet.cell(2, col).value = value
        return workbook

    def _event(self, date, time, title, tournament="Ла Лига"):
        return {
            "date": date, "time": time, "title": title,
            "sport": "Футбол", "tournament": tournament,
            "channel": "SETANTA SPORTS 1",
            "start_at": date + "T" + time + ":00+05:00",
            "platform_start_at": (
                datetime.fromisoformat(date + "T" + time + ":00") -
                __import__("datetime").timedelta(minutes=10)
            ).isoformat() + "+05:00",
            "end_at": (
                datetime.fromisoformat(date + "T" + time + ":00") +
                __import__("datetime").timedelta(minutes=150)
            ).isoformat() + "+05:00",
        }

    def test_preserve_existing_row_and_sort_new_matches(self):
        workbook = self._source()
        previous = self._event("2026-09-24", "23:45",
                               "Португалия - Уэльс", "Лига наций УЕФА")
        early = self._event("2026-09-23", "15:00", "Команда А - Команда Б")
        future = self._event("2026-09-25", "12:00", "Команда В - Команда Г")
        result = build_working_xlsx(workbook, [future, previous, early])
        sheet = result.active
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet["F2"].value, "Команда А - Команда Б")
        self.assertEqual(sheet["F3"].value, "Португалия - Уэльс")
        self.assertEqual(sheet["F4"].value, "Команда В - Команда Г")
        self.assertEqual(sheet["J3"].value, "APPROVED_LIVE_RU")
        self.assertEqual(sheet["Y3"].value, "Футбол. УЕФА Ұлттар лигасы")
        self.assertEqual(sheet["U2"].value, "")
        self.assertEqual(sheet["Y4"].value, "")
        self.assertEqual(
            [sheet.cell(i, 1).value for i in range(2, 5)],
            [60100, 60000, 59900],
        )
        self.assertEqual(sheet["H2"].value, datetime(2026, 9, 23, 14, 50))
        output = BytesIO()
        result.save(output)
        reloaded = load_workbook(BytesIO(output.getvalue()))
        self.assertEqual(reloaded.active["F3"].value, "Португалия - Уэльс")
        reloaded.close()
        result.close()

    def test_priority_follows_template_scale_not_a_fixed_increment(self):
        workbook = self._source()
        sheet = workbook.active
        for column in range(1, 26):
            sheet.cell(3, column).value = sheet.cell(2, column).value
        sheet.cell(2, 1).value = 50000
        sheet.cell(3, 1).value = 49900
        sheet.cell(3, 6).value = "Отдельный матч"
        sheet.cell(3, 8).value = datetime(2026, 9, 25, 11, 50)
        sheet.cell(3, 9).value = datetime(2026, 9, 25, 14, 30)
        added = self._event("2026-09-26", "12:00", "Следующий матч")
        result = build_working_xlsx(workbook, [added])
        self.assertEqual(
            [result.active.cell(row, 1).value for row in (2, 3, 4)],
            [50000, 49900, 49800],
        )
        self.assertEqual(result.active.cell(4, 6).value, "Следующий матч")
        result.close()

    def test_insert_between_approved_25_and_9_point_gaps(self):
        workbook = self._source()
        sheet = workbook.active
        for row in (3, 4):
            for col in range(1, 26):
                sheet.cell(row, col).value = sheet.cell(2, col).value
        # Approved source priorities need not follow a global increment.
        sheet["A2"], sheet["A3"], sheet["A4"] = 50000, 49975, 49966
        sheet["F2"].value = "Первый утверждённый матч"
        sheet["F3"].value = "Второй утверждённый матч"
        sheet["F4"].value = "Третий утверждённый матч"
        sheet["H2"].value = datetime(2026, 9, 24, 23, 35)
        sheet["H3"].value = datetime(2026, 9, 25, 11, 50)
        sheet["H4"].value = datetime(2026, 9, 26, 11, 50)
        extra_a = self._event("2026-09-25", "05:00", "Вставка А")
        extra_b = self._event("2026-09-26", "05:00", "Вставка Б")
        result = build_working_xlsx(workbook, [extra_b, extra_a])
        pairs = {
            result.active.cell(row, 6).value: result.active.cell(row, 1).value
            for row in range(2, result.active.max_row + 1)
        }
        self.assertEqual(
            [pairs["Первый утверждённый матч"],
             pairs["Второй утверждённый матч"],
             pairs["Третий утверждённый матч"]],
            [50000, 49975, 49966],
        )
        self.assertEqual(pairs["Вставка А"], 49988)
        self.assertEqual(pairs["Вставка Б"], 49971)
        self.assertEqual(len(set(pairs.values())), len(pairs))
        result.close()

    def test_exhausted_integer_gap_requires_approval_not_renumber(self):
        workbook = self._source()
        sheet = workbook.active
        for col in range(1, 26):
            sheet.cell(3, col).value = sheet.cell(2, col).value
        sheet["A2"], sheet["A3"] = 50000, 49999
        sheet["F2"].value = "Первый утверждённый матч"
        sheet["F3"].value = "Второй утверждённый матч"
        sheet["H3"].value = datetime(2026, 9, 26, 11, 50)
        new = self._event("2026-09-25", "12:00", "Вставка")
        with self.assertRaisesRegex(InvalidTemplate, "Недостаточно свободных"):
            build_working_xlsx(workbook, [new])
        # The original workbook has not been renumbered.
        self.assertEqual(sheet["A2"].value, 50000)
        self.assertEqual(sheet["A3"].value, 49999)
        workbook.close()

    def test_translation_from_approved_existing_rows(self):
        workbook = self._source()
        new_match = self._event(
            "2026-09-27", "14:00", "Португалия - Уэльс",
            "Лига наций УЕФА"
        )
        result = build_working_xlsx(workbook, [new_match])
        sheet = result.active
        added = next(row for row in sheet.iter_rows(min_row=2)
                     if row[1].value == "27.09")
        self.assertEqual(added[19].value, "ПОРТУГАЛИЯ")
        self.assertEqual(added[20].value, "ПОРТУГАЛИЯ")
        self.assertEqual(added[21].value, "УЭЛЬС")
        self.assertEqual(added[22].value, "УЭЛЬС")
        self.assertEqual(added[24].value, "Футбол. УЕФА Ұлттар лигасы")
        result.close()

    def test_non_match_team1_and_stage_follow_card_memo(self):
        workbook = self._source()
        event = {
            "date": "2026-09-25", "time": "17:00",
            "title": "Гран-при Италии - Квалификация",
            "sport": "Формула-1", "tournament": "",
            "channel": "SETANTA SPORTS 1",
            "start_at": "2026-09-25T17:00:00+05:00",
            "platform_start_at": "2026-09-25T16:50:00+05:00",
            "end_at": "2026-09-25T18:10:00+05:00",
        }
        result = build_working_xlsx(workbook, [event])
        row = next(
            r for r in range(2, result.active.max_row + 1)
            if result.active.cell(r, 6).value == "Гран-при Италии - Квалификация"
        )
        self.assertEqual(result.active.cell(row, 20).value, "ГП ИТАЛИИ")
        self.assertEqual(result.active.cell(row, 22).value, "")
        self.assertEqual(
            result.active.cell(row, 24).value,
            "Формула 1. Квалификация",
        )
        result.close()

    def test_combat_pair_and_tennis_stage_are_editorialized(self):
        workbook = self._source()
        combat = self._event(
            "2026-09-25", "19:00",
            "UFC 332: Силва - Ван - Main Card", "UFC 332",
        )
        combat["sport"] = "ММА"
        tennis = self._event(
            "2026-09-26", "15:00",
            "ATP 250 Ханчжоу: Полуфинал 1", "ATP 250 Ханчжоу",
        )
        tennis["sport"] = "Теннис"
        result = build_working_xlsx(workbook, [combat, tennis])
        values = {
            result.active.cell(r, 6).value: (
                result.active.cell(r, 20).value,
                result.active.cell(r, 22).value,
                result.active.cell(r, 24).value,
            )
            for r in range(2, result.active.max_row + 1)
        }
        self.assertEqual(
            values["UFC 332: Силва - Ван - Main Card"],
            ("UFC 332", "", "ММА. Силва - Ван. Основной кард"),
        )
        self.assertEqual(
            values["ATP 250 Ханчжоу: Полуфинал 1"],
            ("ATP 250 ХАНЧЖОУ", "", "Теннис. ATP 250 Ханчжоу: Полуфинал 1"),
        )
        result.close()

    def test_formula_one_matches_approved_team_and_subtitle_pattern(self):
        workbook = self._source()
        event = {
            "date": "2026-10-03", "time": "13:00",
            "title": "Квалификация", "sport": "Автоспорт",
            "tournament": "Формула-1. Гран-при Бахрейна",
            "channel": "SETANTA SPORTS 1",
            "start_at": "2026-10-03T13:00:00+05:00",
            "platform_start_at": "2026-10-03T12:50:00+05:00",
            "end_at": "2026-10-03T15:10:00+05:00",
        }
        result = build_working_xlsx(workbook, [event])
        row = next(r for r in range(2, result.active.max_row + 1)
                   if result.active.cell(r, 6).value == "Квалификация")
        self.assertEqual(result.active.cell(row, 20).value, "ГП БАХРЕЙНА")
        self.assertEqual(result.active.cell(row, 22).value, "")
        self.assertEqual(result.active.cell(row, 24).value,
                         "Формула 1. Квалификация")
        result.close()

    def test_optional_local_editor_fills_only_missing_kazakh_fields(self):
        from services import editorial_export as exporter
        workbook = self._source()
        event = self._event("2026-09-26", "15:00", "Команда А - Команда Б")
        with patch.dict("os.environ", {"SPORT_AI_EDITOR_AUTO_TRANSLATE": "true"}), \
             patch.object(exporter.ai_pipeline, "editor_suggestions", return_value={
                 "team1_kz": "А КОМАНДАСЫ", "team2_kz": "Б КОМАНДАСЫ",
                 "subtitle_kz": "Футбол. Ла Лига", "time": "00:00",
                 "channel": "WRONG",
             }) as editor:
            result = build_working_xlsx(workbook, [event])
        self.assertEqual(editor.call_count, 1)
        row = result.active.max_row
        self.assertEqual(result.active.cell(row, 7).value, "SETANTA SPORTS 1")
        self.assertEqual(result.active.cell(row, 3).value, "15:00")
        self.assertEqual(result.active.cell(row, 21).value, "А КОМАНДАСЫ")
        self.assertEqual(result.active.cell(row, 23).value, "Б КОМАНДАСЫ")
        self.assertEqual(result.active.cell(row, 25).value, "Футбол. Ла Лига")
        result.close()

    def test_wrong_header_fails_without_losing_data(self):
        workbook = self._source()
        workbook.active["T1"] = "WRONG TEAM 1"
        with self.assertRaises(InvalidTemplate):
            build_working_xlsx(workbook, [])
        self.assertEqual(workbook.active["F2"].value, "Португалия - Уэльс")
        workbook.close()


if __name__ == "__main__":
    unittest.main()
