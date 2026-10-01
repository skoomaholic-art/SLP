"""Offline regressions for two actual provider-style XLSX layouts.

Synthetic schedules only: no production mail, supplier bytes or tokens.
"""
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook

from services.epg_excel import (
    InvalidEPG, import_parsed_epg, parse_supported_epg_channels,
)
from storage.database import SLPDatabase


def xlsx(book):
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def sportplus_sheet():
    book = Workbook()
    sheet = book.active
    sheet["C1"] = '"SPORT PLUS QAZAQSTAN" АРНАСЫНЫҢ ТЕЛЕБАҒДАРЛАМАСЫ'
    sheet["C2"] = "СӘРСЕНБІ, 30 ҚЫРКҮЙЕК"
    sheet["B2"] = "Уақыт"
    sheet["B3"] = .2916666667
    sheet["C3"] = "ҚР Әнұраны"
    sheet["B4"] = .5
    sheet["C4"] = "ММА. ALASH PRIDE 131. ПРЯМАЯ ТРАНСЛЯЦИЯ ИЗ ШЫМКЕНТА"
    sheet["D4"] = .125
    sheet["B5"] = .625
    sheet["C5"] = "ФУТБОЛ ПЛЮС. ПРЯМОЙ ЭФИР"
    sheet["B6"] = .75
    sheet["C6"] = "КПЛ. ТОБОЛ - ЕЛИМАЙ"
    sheet["C7"] = "БЕЙСЕНБІ, 01 ҚАЗАН"
    sheet["B9"] = .85
    sheet["C9"] = "БӘЙГЕ. АТ ШАБЫСЫ. АЛМАТЫДАН ТІКЕЛЕЙ ЭФИР"
    sheet["B8"] = .7
    sheet["C8"] = "ФУТЗАЛ. КОМАНДА А - КОМАНДА Б. ПРЯМАЯ ТРАНСЛЯЦИЯ"
    sheet["D8"] = .05
    return book


def qazsport_sheet():
    book = Workbook()
    sheet = book.active
    sheet["C8"] = "«QAZSPORT» ТЕЛЕАРНАСЫ 28.09.2026 - 04.10.2026"
    sheet["A9"] = datetime(2026, 9, 29)
    sheet["B9"] = .5
    sheet["C9"] = "ФУТБОЛ. Лига. Команда А - Команда Б. Тікелей эфир"
    sheet["D9"] = .10
    sheet["F9"] = "трансляция"
    sheet["A10"] = datetime(2026, 9, 29)
    sheet["B10"] = .6
    sheet["C10"] = "ХОККЕЙ. КХЛ. «Барыс» - «Ак Барс». Тікелей эфир"
    sheet["A11"] = datetime(2026, 9, 29)
    sheet["B11"] = .7
    sheet["C11"] = "ФУТБОЛ. Студиялық бағдарлама. Тікелей эфир"
    sheet["F11"] = "трансляция"
    sheet["A12"] = datetime(2026, 9, 29)
    sheet["B12"] = .8
    sheet["C12"] = "ДЗЮДО. Финал"
    sheet["F12"] = "трансляция"
    sheet["A13"] = datetime(2026, 9, 29)
    sheet["B13"] = .99
    sheet["C13"] = "ДЗЮДО. Финал. Тікелей эфир"
    sheet["D13"] = .08
    sheet["A14"] = datetime(2026, 9, 29)
    sheet["B14"] = 1.05
    sheet["C14"] = "Бокс. Финал. Тікелей эфир"
    sheet["D14"] = .07
    return book


class OfficialSupplierEPGTests(unittest.TestCase):
    def test_sportplus_uses_explicit_live_and_kazakh_dates(self):
        content = xlsx(sportplus_sheet())
        parsed, = parse_supported_epg_channels(
            content, "сетка Канала.xlsx", today=date(2026, 9, 29),
            context="Письмо SPORT PLUS QAZAQSTAN",
        )
        self.assertEqual(parsed.channel, "SPORT+ Qazaqstan")
        self.assertEqual(
            [(e["date"], e["time"], e["sport"]) for e in parsed.events],
            [
                ("2026-09-30", "12:00", "ММА"),
                ("2026-10-01", "16:48", "Футзал"),
                ("2026-10-01", "20:24", "Конный спорт"),
            ],
        )
        self.assertEqual(parsed.events[0]["estimated_broadcast_end"], "15:00")
        self.assertEqual(parsed.events[0]["end_estimation_method"], "provider_duration")
        self.assertEqual(parsed.all_programmes, 6)
        self.assertNotIn("ФУТБОЛ ПЛЮС", str(parsed.events))
        self.assertNotIn("ТОБОЛ", str(parsed.events))
        with tempfile.TemporaryDirectory() as dirname:
            database = SLPDatabase(Path(dirname) / "events.db")
            self.assertEqual(import_parsed_epg(database, parsed)["status"], "imported")
            self.assertEqual(len(database.load_active_source_snapshot(
                "email_epg_sportplus", "2026-09-30",
            )), 1)
            self.assertEqual(import_parsed_epg(database, parsed)["status"], "already_imported")

    def test_qazsport_avoids_barys_studio_and_unmarked_replay(self):
        parsed, = parse_supported_epg_channels(
            xlsx(qazsport_sheet()),
            "28.09-04.10.2026 Программа телеканала QAZSPORT.xlsx",
            today=date(2026, 9, 29),
        )
        self.assertEqual(parsed.channel, "QAZSPORT HD")
        self.assertEqual(
            [(e["date"], e["time"], e["sport"]) for e in parsed.events],
            [
                ("2026-09-29", "12:00", "Футбол"),
                ("2026-09-29", "23:46", "Дзюдо"),
                ("2026-09-30", "01:12", "Бокс"),
            ],
        )
        self.assertNotIn("Барыс", str(parsed.events))
        self.assertNotIn("Студий", str(parsed.events))
        self.assertEqual(parsed.events[-1]["estimated_broadcast_end_date"],
                         "2026-09-30")

    def test_identity_precedes_contradictory_filename(self):
        with self.assertRaisesRegex(InvalidEPG, "противоречит"):
            parse_supported_epg_channels(
                xlsx(qazsport_sheet()), "EPG Setanta Sports 1 Kazakhstan.xlsx",
                today=date(2026, 9, 29),
            )

    def test_official_format_refuses_unsupported_old_week_without_year(self):
        with self.assertRaisesRegex(InvalidEPG, "Год недели"):
            parse_supported_epg_channels(
                xlsx(sportplus_sheet()), "сетка Канала.xlsx",
                today=date(2027, 8, 1),
            )

    def test_unknown_sender_context_cannot_invent_a_channel(self):
        book = Workbook()
        book.active["A1"] = "LIVE. Футбол. Команда А - Команда Б"
        with self.assertRaises(InvalidEPG):
            parse_supported_epg_channels(
                xlsx(book), "сетка Канала.xlsx",
                today=date(2026, 9, 29),
                context="QAZSPORT SPORT+ Q LEAGUE",
            )


if __name__ == "__main__":
    unittest.main()
