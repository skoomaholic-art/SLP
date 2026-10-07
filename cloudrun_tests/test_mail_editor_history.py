"""Offline Gmail and editorial regressions; never contact Gmail or send mail."""
import base64
from datetime import date
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from cryptography.fernet import Fernet
from openpyxl import Workbook

from services import gmail_integration as gmail, editorial_store as edits
from storage.database import SLPDatabase
from cloudrun_web import event_rows


def sample_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet["C3"] = "30 сентября"
    sheet["B4"] = "AST"
    sheet["B5"] = 0.50
    sheet["C5"] = "LIVE. Футбол. АПЛ, 6 тур, Команда А - Команда Б"
    sheet["B6"] = 0.57
    sheet["C6"] = "Повтор матча"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


class GmailOfflineTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.db = SLPDatabase(Path(self.temporary.name) / "storage.db")
        gmail.init_gmail_schema(self.db)

    def test_oauth_requires_owner_identity_and_encrypts_refresh_token(self):
        secrets = {
            "SPORT_GMAIL_CLIENT_ID": "client-123",
            "SPORT_GMAIL_CLIENT_SECRET": "do-not-print",
            "SPORT_GMAIL_TOKEN_KEY": Fernet.generate_key().decode(),
            "SPORT_PUBLIC_URL": "https://sport.example",
        }
        with patch.dict("os.environ", secrets):
            link = gmail.start_oauth(self.db, username="Skoomaholic")
            self.assertIn("accounts.google.com", link)
            with self.db._connect() as conn:
                state_row = conn.execute(
                    "SELECT state_sha FROM gmail_oauth_states"
                ).fetchone()
            from urllib.parse import parse_qs, urlparse
            state = parse_qs(urlparse(link).query)["state"][0]
            self.assertNotEqual(state_row["state_sha"], state)
            with self.assertRaises(gmail.GmailTransportError):
                gmail.complete_oauth(
                    self.db, username="Дания", state=state, code="code-123"
                )
            # OAuth state is one-use even if a different account tries it.
            with self.assertRaises(gmail.GmailTransportError):
                gmail.complete_oauth(
                    self.db, username="Skoomaholic", state=state, code="code-123"
                )
            link = gmail.start_oauth(self.db, username="Skoomaholic")
            state = parse_qs(urlparse(link).query)["state"][0]
            with patch.object(gmail, "_json_http", return_value={
                "access_token": "fake-access", "refresh_token": "fake-private-token"
            }), patch.object(gmail, "_json_api", return_value={
                "emailAddress": gmail.OWNER_ACCOUNT
            }):
                gmail.complete_oauth(
                    self.db, username="Skoomaholic", state=state, code="code-123"
                )
            with self.db._connect() as conn:
                encrypted = conn.execute(
                    "SELECT refresh_token_cipher FROM gmail_oauth WHERE id=1"
                ).fetchone()[0]
            self.assertNotIn(b"fake-private-token", bytes(encrypted))
            self.assertTrue(gmail.status(self.db)["connected"])

    def test_sync_stages_file_and_change_letter_without_import(self):
        blob = sample_xlsx()
        msg1, msg2 = "abcde1234567", "abcde1234568"
        envelope = {
            "id": msg1, "internalDate": "1790812800000",
            "snippet": "EPG сетка Setanta",
            "payload": {
                "headers": [
                    {"name": "Subject", "value": "EPG Setanta Sports 2 Kazakhstan"},
                    {"name": "From", "value": "Поставщик <supplier@example.test>"},
                ],
                "parts": [{
                    "filename": "EPG Setanta Sports 2 Kazakhstan 29.09.26 - 05.10.26_MEDIA.xlsx",
                    "body": {"size": len(blob), "attachmentId": "attachment-1"},
                }],
            },
        }
        change_letter = {
            "id": msg2, "internalDate": "1790812800000",
            "snippet": "Изменение эфира SPORT PLUS",
            "payload": {
                "headers": [
                    {"name": "Subject", "value": "Изменения SPORT PLUS"},
                    {"name": "From", "value": "Поставщик <supplier@example.test>"},
                ],
            },
        }
        def fake_api(url, token, payload=None):
            if "messages?" in url:
                return {"messages": [{"id": msg1}, {"id": msg2}]}
            if msg1 in url and "/attachments/" in url:
                return {"data": base64.urlsafe_b64encode(blob).decode().rstrip("=")}
            if msg1 in url:
                return envelope
            if msg2 in url:
                return change_letter
            raise AssertionError(url)
        with patch.dict("os.environ", {"SPORT_GMAIL_AUTO_IMPORT": "false"}), \
             patch.object(gmail, "_access_token", return_value="secret-access"), \
             patch.object(gmail, "_json_api", side_effect=fake_api):
            first = gmail.sync_inbox(self.db)
            second = gmail.sync_inbox(self.db)
        self.assertEqual(first["new_attachments"], 1)
        self.assertEqual(second["new_attachments"], 0)
        self.assertEqual(self.db.active_event_count(), 0)
        notices = gmail.list_notices(self.db)
        self.assertEqual(len(notices), 2)
        pending = next(x for x in notices if x["status"] == "pending")
        review = next(x for x in notices if x["status"] == "review")
        self.assertEqual(pending["detected_channel"], "SETANTA SPORTS 2")
        self.assertFalse(review["has_attachment"])
        result = gmail.approve_notice(self.db, pending["id"], username="Вадим")
        self.assertEqual(result["status"], "imported")
        self.assertGreater(self.db.active_event_count(), 0)
        self.assertEqual(gmail.approve_notice.__name__, "approve_notice")
        self.assertFalse(next(x for x in gmail.list_notices(self.db)
                              if x["id"] == pending["id"])["has_attachment"])

    def test_older_pending_notice_closes_as_superseded_after_newer_import(self):
        def workbook_at(excel_time):
            workbook = Workbook()
            sheet = workbook.active
            sheet["C3"] = "30 сентября"
            sheet["B4"] = "AST"
            sheet["B5"] = excel_time
            sheet["C5"] = "LIVE. Футбол. АПЛ, 6 тур, Команда А - Команда Б"
            stream = BytesIO()
            workbook.save(stream)
            workbook.close()
            return stream.getvalue()

        old_blob = workbook_at(0.50)
        new_blob = workbook_at(0.55)
        old_id, new_id = "abcde1234581", "abcde1234582"

        def envelope(message_id, attachment_id, blob, millis):
            return {
                "id": message_id,
                "internalDate": str(millis),
                "snippet": "EPG Setanta",
                "payload": {
                    "headers": [
                        {"name": "Subject", "value": "EPG Setanta Sports 2 Kazakhstan"},
                        {"name": "From", "value": "Supplier <supplier@example.test>"},
                    ],
                    "parts": [{
                        "filename": "EPG Setanta Sports 2 Kazakhstan 29.09.26 - 05.10.26_MEDIA.xlsx",
                        "body": {"size": len(blob), "attachmentId": attachment_id},
                    }],
                },
            }

        messages = {
            old_id: envelope(old_id, "old-grid", old_blob, 1790812800000),
            new_id: envelope(new_id, "new-grid", new_blob, 1790816400000),
        }
        attachments = {"old-grid": old_blob, "new-grid": new_blob}

        def fake_api(url, token, payload=None):
            if "messages?" in url:
                return {"messages": [{"id": old_id}, {"id": new_id}]}
            if "/attachments/" in url:
                for attachment_id, blob in attachments.items():
                    if attachment_id in url:
                        return {
                            "data": base64.urlsafe_b64encode(blob).decode().rstrip("=")
                        }
            for message_id, value in messages.items():
                if message_id in url:
                    return value
            raise AssertionError(url)

        with patch.dict("os.environ", {"SPORT_GMAIL_AUTO_IMPORT": "false"}), \
             patch.object(gmail, "_access_token", return_value="token"), \
             patch.object(gmail, "_json_api", side_effect=fake_api):
            gmail.sync_inbox(self.db)

        notices = [
            row for row in gmail.list_notices(self.db)
            if row["status"] == "pending"
        ]
        self.assertEqual(len(notices), 2)
        newer = max(notices, key=lambda row: row["received_at"])
        older = min(notices, key=lambda row: row["received_at"])
        applied = gmail.approve_notice(self.db, newer["id"], username="Editor")
        self.assertEqual(applied["status"], "imported")

        stale = gmail.approve_notice(self.db, older["id"], username="Editor")
        self.assertEqual(stale["status"], "superseded")
        self.assertFalse(stale["applied"])
        stored = next(
            row for row in gmail.list_notices(self.db)
            if row["id"] == older["id"]
        )
        self.assertEqual(stored["status"], "superseded")
        self.assertFalse(stored["has_attachment"])

    def test_checkpoint_waits_for_retryable_message_then_advances(self):
        good_id, bad_id = "abcde1234501", "abcde1234502"
        good = {
            "id": good_id, "internalDate": "1790812800000", "snippet": "hello",
            "payload": {"headers": [
                {"name": "Subject", "value": "Other"},
                {"name": "From", "value": "person@example.test"},
            ]},
        }

        def fake_api(url, token, payload=None):
            if "messages?" in url:
                return {"messages": [{"id": good_id}, {"id": bad_id}]}
            if bad_id in url:
                raise gmail.GmailTransportError("temporary")
            if good_id in url:
                return good
            raise AssertionError(url)

        with patch.object(gmail, "_access_token", return_value="token"), \
             patch.object(gmail, "_json_api", side_effect=fake_api):
            first = gmail.sync_inbox(self.db)
            second = gmail.sync_inbox(self.db)
            third = gmail.sync_inbox(self.db)
        self.assertFalse(first["checkpoint_advanced"])
        self.assertFalse(second["checkpoint_advanced"])
        self.assertTrue(third["checkpoint_advanced"])
        self.assertEqual(third["failed_messages"], 1)
        self.assertEqual(third["quarantined_messages"], 1)
        with self.db._connect() as conn:
            state = conn.execute(
                "SELECT last_internal_date_ms FROM gmail_sync_state WHERE id=1"
            ).fetchone()
            error = conn.execute(
                "SELECT attempts,resolved_at FROM gmail_processing_errors "
                "WHERE message_id=?", (bad_id,),
            ).fetchone()
        self.assertEqual(state["last_internal_date_ms"], 1790812800000)
        self.assertEqual(error["attempts"], 3)
        self.assertEqual(error["resolved_at"], "")
        self.assertTrue(all(query.startswith("after:")
                            for query in gmail._search_queries(self.db)))

    def test_approved_structure_is_reused_for_generic_filename(self):
        blob = sample_xlsx()
        first_id, second_id = "abcde1234511", "abcde1234512"

        def envelope(message_id, filename, attachment_id, millis):
            return {
                "id": message_id, "internalDate": str(millis),
                "snippet": "актуальная сетка", "payload": {
                    "headers": [
                        {"name": "Subject", "value": "Расписание"},
                        {"name": "From", "value": "Supplier <same@example.test>"},
                    ],
                    "parts": [{"filename": filename, "body": {
                        "size": len(blob), "attachmentId": attachment_id,
                    }}],
                },
            }

        messages = {first_id: envelope(
            first_id, "EPG Setanta Sports 2 Kazakhstan.xlsx", "first-file",
            1790812800000,
        )}

        def fake_api(url, token, payload=None):
            if "messages?" in url:
                return {"messages": [{"id": key} for key in messages]}
            if "/attachments/" in url:
                return {"data": base64.urlsafe_b64encode(blob).decode().rstrip("=")}
            for key, value in messages.items():
                if key in url:
                    return value
            raise AssertionError(url)

        with patch.dict("os.environ", {"SPORT_GMAIL_AUTO_IMPORT": "false"}), \
             patch.object(gmail, "_access_token", return_value="token"), \
             patch.object(gmail, "_json_api", side_effect=fake_api):
            gmail.sync_inbox(self.db)
            first_notice = gmail.list_notices(self.db)[0]
            gmail.approve_notice(self.db, first_notice["id"], username="Editor")
            messages[second_id] = envelope(
                second_id, "сетка Канала.xlsx", "second-file", 1790816400000,
            )
            outcome = gmail.sync_inbox(self.db)
        self.assertEqual(outcome["new_attachments"], 1)
        second_notice = next(
            row for row in gmail.list_notices(self.db)
            if row["source_document_sha256"] and row["status"] == "pending"
        )
        self.assertEqual(second_notice["detected_channel"], "SETANTA SPORTS 2")
        self.assertEqual(second_notice["channel_detection_method"], "confirmed_format")

    def test_forwarded_mojibake_is_repaired_before_change_detection(self):
        broken = "РќР° 6 РѕРєС‚СЏР±СЂСЏ РґРѕР±Р°РІРёР»Рё РџСЂСЏРјРѕР№ СЌС„РёСЂ РўРµРЅРЅРёСЃР°"
        repaired = gmail._repair_forwarded_text(broken)
        self.assertIn("На 6 октября", repaired)
        self.assertIn("Прямой эфир Тенниса", repaired)

    def test_forwarded_generic_attachment_uses_body_context_and_workbook(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet["B1"] = "ПРОГРАММА ПЕРЕДАЧ ТЕЛЕКАНАЛА SETANTA SPORTS 1 KAZAKHSTAN"
        sheet["B4"] = "30 сентября"
        sheet["A5"] = "AST"
        sheet["A6"] = 0.5
        sheet["B6"] = "LIVE. Футбол. АПЛ, 6 тур, Команда А - Команда Б"
        stream = BytesIO(); workbook.save(stream); workbook.close()
        blob = stream.getvalue()
        msg = "abcde1234570"
        forwarded_text = (
            "От: Sabina Nazarova <sabina@silkwaymedia.com> "
            "Тема: актуальная сетка Setanta. Со вторника есть изменения."
        )
        body = base64.urlsafe_b64encode(forwarded_text.encode()).decode().rstrip("=")
        envelope = {
            "id": msg, "internalDate": "1790812800000",
            "snippet": "Пересылаю сетку",
            "payload": {
                "headers": [
                    {"name": "Subject", "value": "FW: актуальная сетка"},
                    {"name": "From", "value": "Alexandr <alexandr.petrossov@fmedia.kz>"},
                ],
                "parts": [
                    {"mimeType": "text/plain", "body": {"data": body}},
                    {
                        "filename": "сетка Канала.xlsx",
                        "body": {"size": len(blob), "attachmentId": "attachment-generic"},
                    },
                ],
            },
        }

        def fake_api(url, token, payload=None):
            if "messages?" in url:
                return {"messages": [{"id": msg}]}
            if "/attachments/" in url:
                return {"data": base64.urlsafe_b64encode(blob).decode().rstrip("=")}
            if msg in url:
                return envelope
            raise AssertionError(url)

        with patch.dict("os.environ", {"SPORT_GMAIL_AUTO_IMPORT": "false"}), \
             patch.object(gmail, "_access_token", return_value="secret-access"), \
             patch.object(gmail, "_json_api", side_effect=fake_api):
            outcome = gmail.sync_inbox(self.db)

        self.assertEqual(outcome["new_attachments"], 1)
        notice = gmail.list_notices(self.db)[0]
        self.assertEqual(notice["detected_channel"], "SETANTA SPORTS 1")
        self.assertEqual(notice["filename"], "сетка Канала.xlsx")
        self.assertEqual(self.db.active_event_count(), 0)

    def test_verified_new_epg_imports_automatically_and_only_once(self):
        blob = sample_xlsx()
        msg_id = "abcde1234571"
        envelope = {
            "id": msg_id, "internalDate": "1790812800000",
            "snippet": "EPG Setanta",
            "payload": {
                "headers": [
                    {"name": "Subject", "value": "Setanta Sports 2 Kazakhstan EPG"},
                    {"name": "From", "value": "Supplier <supplier@example.test>"},
                ],
                "parts": [{
                    "filename": "EPG Setanta Sports 2 Kazakhstan 29.09.26 - 05.10.26_MEDIA.xlsx",
                    "body": {"size": len(blob), "attachmentId": "attachment-auto"},
                }],
            },
        }
        def fake_api(url, token, payload=None):
            if "messages?" in url:
                return {"messages": [{"id": msg_id}]}
            if "/attachments/" in url:
                return {"data": base64.urlsafe_b64encode(blob).decode().rstrip("=")}
            if msg_id in url:
                return envelope
            raise AssertionError(url)
        with patch.dict("os.environ", {"SPORT_GMAIL_AUTO_IMPORT": "true"}), \
             patch.object(gmail, "_access_token", return_value="secret-access"), \
             patch.object(gmail, "_json_api", side_effect=fake_api):
            first = gmail.sync_inbox(self.db)
            second = gmail.sync_inbox(self.db)
        self.assertEqual(first["auto_imported"], 1)
        self.assertEqual(second["auto_imported"], 0)
        self.assertEqual(self.db.active_event_count(), 1)
        notices = gmail.list_notices(self.db)
        self.assertEqual(notices[0]["status"], "imported")
        self.assertFalse(notices[0]["has_attachment"])

    def test_suspicious_update_with_no_live_does_not_erase_previous_events(self):
        original = sample_xlsx()
        from services.epg_excel import import_parsed_epg, parse_epg_xlsx
        imported = parse_epg_xlsx(
            original,
            "EPG Setanta Sports 2 Kazakhstan 29.09.26 - 05.10.26_MEDIA.xlsx",
            today=date(2026, 9, 29),
        )
        import_parsed_epg(self.db, imported)
        result = gmail._safe_auto_apply(self.db, 123, type(imported)(
            channel=imported.channel, filename=imported.filename,
            content_hash="another-file", scope_dates=imported.scope_dates,
            all_programmes=imported.all_programmes, events=(),
        ))
        self.assertFalse(result[0])
        self.assertEqual(self.db.active_event_count(), 1)

    def test_two_channels_inside_one_mail_attachment_are_imported_once_each(self):
        wb = Workbook()
        for channel in ("Q ARENA", "Q LEAGUE"):
            sheet = wb.active if channel == "Q ARENA" else wb.create_sheet()
            sheet.title = channel
            sheet["C3"] = "29 сентября"
            sheet["B4"] = "AST"
            sheet["B5"] = .5
            sheet["C5"] = (
                "LIVE. Футбол. КПЛ, " +
                ("А - Б" if channel == "Q ARENA" else "В - Г")
            )
            sheet["B6"] = .6
            sheet["C6"] = "Новости"
        stream = BytesIO()
        wb.save(stream)
        wb.close()
        blob = stream.getvalue()
        msg = "abcde1234588"
        envelope = {
            "id": msg, "internalDate": "1790812800000",
            "snippet": "Сетка QSport", "payload": {
                "headers": [
                    {"name": "Subject", "value": "EPG QSport"},
                    {"name": "From", "value": "supplier@example.test"},
                ],
                "parts": [{
                    "filename": "сетка Канала.xlsx",
                    "body": {"size": len(blob), "attachmentId": "multi-file"},
                }],
            },
        }
        def fake_api(url, token, payload=None):
            if "messages?" in url:
                return {"messages": [{"id": msg}]}
            if "/attachments/" in url:
                return {"data": base64.urlsafe_b64encode(blob).decode().rstrip("=")}
            if msg in url:
                return envelope
            raise AssertionError(url)
        with patch.dict("os.environ", {"SPORT_GMAIL_AUTO_IMPORT": "true"}), \
             patch.object(gmail, "_access_token", return_value="fake-token"), \
             patch.object(gmail, "_json_api", side_effect=fake_api):
            first = gmail.sync_inbox(self.db)
            again = gmail.sync_inbox(self.db)
        self.assertEqual(first["auto_imported"], 2)
        self.assertEqual(first["new_attachments"], 2)
        self.assertEqual(again["new_attachments"], 0)
        self.assertEqual(self.db.active_event_count(), 2)
        self.assertEqual(
            {notice["detected_channel"] for notice in gmail.list_notices(self.db)},
            {"Q ARENA", "Q LEAGUE"},
        )

    def test_test_mail_does_not_send_unless_explicitly_enabled(self):
        with patch.dict("os.environ", {"SPORT_GMAIL_ENABLE_TEST_SEND": "false"}):
            with patch.object(gmail, "_json_api") as post:
                with self.assertRaises(gmail.GmailNotConfigured):
                    gmail.test_request(
                        self.db, username="Дания", category="q",
                        destination="alexandr.petrossov@fmedia.kz"
                    )
                post.assert_not_called()

    def test_request_records_period_and_rejects_duplicate(self):
        sent = {"id": "abcde1234599", "threadId": "thread-1"}
        with patch.dict("os.environ", {
            "SPORT_GMAIL_ENABLE_TEST_SEND": "true",
            "SPORT_GMAIL_MAIL_MODE": "test",
            "SPORT_MAIL_TEST_TO": "alexandr.petrossov@fmedia.kz",
        }), patch.object(gmail, "_access_token", return_value="token"), \
             patch.object(gmail, "_json_api", return_value=sent):
            result = gmail.send_schedule_request(
                self.db, username="Editor", category="q",
                period_start="2026-10-01", period_end="2026-10-07",
                destination="alexandr.petrossov@fmedia.kz",
            )
            self.assertEqual(result["mode"], "test")
            with self.assertRaisesRegex(gmail.GmailTransportError, "уже ожидает"):
                gmail.send_schedule_request(
                    self.db, username="Editor", category="q",
                    period_start="2026-10-01", period_end="2026-10-07",
                    destination="alexandr.petrossov@fmedia.kz",
                )
        requests = gmail.list_requests(self.db)
        self.assertEqual(requests[0]["period_start"], "2026-10-01")
        self.assertEqual(requests[0]["requested_channels"], [
            "Q LEAGUE", "Q ARENA", "Q FOOTBALL",
        ])

    def test_production_request_needs_configured_supplier_recipient(self):
        with patch.dict("os.environ", {
            "SPORT_GMAIL_MAIL_MODE": "production",
            "SPORT_MAIL_QSPORT_TO": "",
            "SPORT_MAIL_SETANTA_TO": "",
        }):
            with self.assertRaisesRegex(gmail.GmailNotConfigured, "получатель"):
                gmail.send_schedule_request(
                    self.db, username="Editor", category="q",
                    period_start="2026-10-01", period_end="2026-10-07",
                )


class HistoryTests(unittest.TestCase):
    def test_history_is_persistent_and_editor_overrides_do_not_destroy_source(self):
        with tempfile.TemporaryDirectory() as folder:
            db = SLPDatabase(Path(folder) / "events.db")
            event = {
                "source": "tvguide", "date": "2026-10-03", "time": "12:00",
                "sport": "Футбол", "tournament": "Лига",
                "title": "Команда А - Команда Б",
                "raw_title": "Команда А - Команда Б",
                "channel": "SETANTA SPORTS 2", "is_live": True,
                "is_live_broadcast": True,
            }
            db.upsert_source_snapshot(
                run_id="r1", source="tvguide", scope_date="2026-10-03",
                events=[event]
            )
            amended = {**event, "tournament": "Лига, 2-й тур"}
            db.upsert_source_snapshot(
                run_id="r2", source="tvguide", scope_date="2026-10-03",
                events=[amended]
            )
            storage_id = db.event_revisions(limit=3)[0]["storage_id"]
            self.assertEqual(
                [x["change_kind"] for x in db.event_revisions(storage_id)],
                ["source_changed", "added_to_source"]
            )
            edits.apply_edit(
                db, storage_id=storage_id, username="Дания",
                values={"title": "Команда А - Команда В",
                        "team1_kz": "А командасы", "time": "12:30"}
            )
            visible = event_rows(db, "2026-10-03", "2026-10-03")
            self.assertEqual(len(visible), 1)
            self.assertEqual(visible[0]["title"], "Команда А - Команда В")
            self.assertEqual(visible[0]["team1_kz"], "А командасы")
            self.assertEqual(visible[0]["time"], "12:30")
            self.assertEqual(len(edits.edit_history(db, storage_id=storage_id)), 1)
            db.upsert_source_snapshot(
                run_id="r3", source="tvguide", scope_date="2026-10-03",
                events=[]
            )
            history = db.event_revisions(storage_id)
            self.assertEqual(history[0]["change_kind"], "removed_from_source")
            self.assertEqual(len(history), 3)
            with db._connect() as conn:
                row = conn.execute(
                    "SELECT payload_json FROM events WHERE storage_id=?",
                    (storage_id,),
                ).fetchone()
            self.assertIsNotNone(row["payload_json"])


if __name__ == "__main__":
    unittest.main()
