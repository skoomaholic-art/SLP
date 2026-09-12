import unittest
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from services.export_xlsx import (
    EXPORT_MODE_FINISH_UPCOMING,
    EXPORT_MODE_LIVE,
    build_schedule_xlsx,
)


KZ = ZoneInfo("Asia/Almaty")


def make_event(time_text: str, end_text: str, title: str, *, direct: bool) -> dict:
    return {
        "source": "tvguide",
        "source_url": "https://example.test/guide",
        "channel": "Setanta Sports 1",
        "date": "2026-09-12",
        "time": time_text,
        "timezone": "Asia/Almaty",
        "sport": "Футбол",
        "tournament": "Test League",
        "title": title,
        "is_live": direct,
        "is_live_broadcast": direct,
        "raw_title": title,
        "estimated_broadcast_end_date": "2026-09-12",
        "estimated_broadcast_end": end_text,
        "end_estimation_method": "provider_epg",
        "end_confidence": "high",
    }


class ExportXlsxTests(unittest.TestCase):
    def setUp(self):
        self.now = datetime(2026, 9, 12, 22, 28, tzinfo=KZ)
        self.events = [
            make_event("16:00", "18:00", "Finished", direct=False),
            make_event("21:30", "23:30", "On air", direct=True),
            make_event("23:30", "23:59", "Upcoming", direct=False),
        ]

    def test_live_export_contains_only_temporally_live_rows(self):
        payload = build_schedule_xlsx(
            self.events,
            mode=EXPORT_MODE_LIVE,
            now=self.now,
        )
        workbook = load_workbook(BytesIO(payload))
        sheet = workbook["Расписание"]
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["G2"].value, "On air")
        self.assertEqual(sheet["D2"].value, "LIVE")
        self.assertEqual(sheet["L2"].value, "Да")
        self.assertEqual(sheet["M2"].value, "Asia/Almaty")

    def test_finish_upcoming_export_excludes_live_rows(self):
        payload = build_schedule_xlsx(
            self.events,
            mode=EXPORT_MODE_FINISH_UPCOMING,
            now=self.now,
        )
        workbook = load_workbook(BytesIO(payload))
        sheet = workbook["Расписание"]
        self.assertEqual(sheet.max_row, 3)
        statuses = {sheet[f"D{row}"].value for row in range(2, 4)}
        self.assertEqual(statuses, {"FINISHED", "UPCOMING"})
        titles = {sheet[f"G{row}"].value for row in range(2, 4)}
        self.assertNotIn("On air", titles)

    def test_export_uses_distinct_temporal_and_direct_fields(self):
        payload = build_schedule_xlsx(self.events, now=self.now)
        workbook = load_workbook(BytesIO(payload))
        sheet = workbook["Расписание"]
        self.assertEqual(sheet["D1"].value, "Статус эфира")
        self.assertEqual(sheet["L1"].value, "Прямой эфир")
        self.assertEqual(sheet["M1"].value, "Часовой пояс")
        self.assertEqual(sheet["L2"].value, "Нет")
        self.assertEqual(sheet["L3"].value, "Да")


if __name__ == "__main__":
    unittest.main()
